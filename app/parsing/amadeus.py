"""Parser for the Amadeus daily sales report (.xlsx).

The eight required Amadeus columns and the optional ones (TAX, COMM, FP, CONFIRMED) are read,
located by header name. Anything else in the sheet (hand-typed columns, stray formulas) has no
effect on the result, except that unknown header titles are named in one warning.

The workbook is streamed (``read_only``) so a tiny file that declares a cell on row
1,048,576 cannot make us materialise a million empty rows, and the archive is checked
for zip bombs before openpyxl sees it. XML goes through defusedxml (no DTDs/entities).
"""

from __future__ import annotations

import dataclasses
import io
import itertools
import re
import unicodedata
import zipfile
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any, BinaryIO, Literal

import openpyxl
import openpyxl.xml
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 30 * 1024 * 1024
MAX_ENTRY_BYTES = 20 * 1024 * 1024
MAX_ZIP_ENTRIES = 1000
MAX_DATA_ROWS = 50_000  # non-blank rows below the header
MAX_COLUMNS = 30  # only this many leading columns are ever read
HEADER_SCAN_ROWS = 20
MAX_WARNINGS = 100
MAX_UNKNOWN_COLUMNS = 10  # header titles named in the "unknown columns" warning
UNKNOWN_COLUMN_CHARS = 30
WARNING_VALUE_CHARS = 40

PaxType = Literal["ADT", "CHD", "INF"]

FLAG_MISSING_COST = "missing_cost"
FLAG_UNKNOWN_TRNC = "unknown_trnc"
FLAG_MISSING_BOOKING_REF = "missing_booking_ref"
FLAG_MISSING_STAFF_SIGN = "missing_staff_sign"
FLAG_DUPLICATE_IN_FILE = "duplicate_in_file"
FLAG_FIELD_TOO_LONG = "field_too_long"
FLAG_COST_ROUNDED = "cost_rounded"
FLAG_CANCELLED = "cancelled"

TRNC_ISSUED = "TKTT"
# A cancelled document: stored like any line, but never a sale (the accountant deletes these
# lines before working on a day), so it is left out of every count, total and price question.
TRNC_CANCELLED = "CANX"

MAX_NAME_CHARS = 200  # passenger names are free text; anything longer is cut and flagged
KNOWN_FLAGS = frozenset(
    {
        FLAG_MISSING_COST,
        FLAG_UNKNOWN_TRNC,
        FLAG_MISSING_BOOKING_REF,
        FLAG_MISSING_STAFF_SIGN,
        FLAG_DUPLICATE_IN_FILE,
        FLAG_FIELD_TOO_LONG,
        FLAG_COST_ROUNDED,
        FLAG_CANCELLED,
    }
)

# Column lengths of ``app.models.Ticket`` and ``app.models.AmadeusLine``; longer values are
# truncated and flagged.
FIELD_LIMITS = {
    "seq_no": 32,
    "airline": 8,
    "staff_sign": 32,
    "transaction_type": 16,
    "booking_ref": 16,
    "form_of_payment": 16,
    "confirmed": 32,
}
COST_LIMIT = Decimal(10) ** 12  # Numeric(14, 2)
_CENT = Decimal("0.01")

_MONTHS = {
    m: i
    for i, m in enumerate(
        ["JAN", "FEB", "MAR", "APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC"],
        start=1,
    )
}
_RANGE_RE = re.compile(r"^(\d{1,2})([A-Za-z]{3})$")
_CURRENT_RE = re.compile(r"^(\d{1,2})-([A-Za-z]{3})-(\d{4})$")
_PAX_SUFFIX_RE = re.compile(r"\s*\((INF|CHD)\)$", re.IGNORECASE)
_PAX_MARKER_RE = re.compile(r"\((INF|CHD)", re.IGNORECASE)
_TICKET_NO_RE = re.compile(r"\d{8,15}", re.ASCII)

# header name -> internal key
_REQUIRED = {
    "SEQ NO": "seq_no",
    "A/L": "airline",
    "DOC NUMBER": "ticket_no",
    "TOTAL DOC": "cost",
    "PAX NAME": "passenger_name",
    "AS": "staff_sign",
    "TRNC": "transaction_type",
    "RLOC": "booking_ref",
}
# Optional columns (the newer 16-column export); absent ones give None.
_OPTIONAL = {
    "TAX": "tax",
    "COMM": "comm",
    "FP": "form_of_payment",
    "CONFIRMED": "confirmed",
}
# Columns that are deliberately not read, so they get no warning: Amadeus columns of the newer
# export (FEE, AGENT, REPORT IND, NDC) and the accountant's own hand-typed columns, which are in
# every working file (P.B, AG, CLT, AP).
_KNOWN_IGNORED = frozenset({"FEE", "AGENT", "REPORT IND", "NDC", "P.B", "AG", "CLT", "AP"})


class ParseError(Exception):
    """The file is not a readable Amadeus daily sales report."""


@dataclass(frozen=True)
class ReportHeader:
    report_range_raw: str
    report_date: date
    generated_on: date


@dataclass(frozen=True)
class ParsedTicket:
    seq_no: str
    airline: str
    ticket_no: str
    cost: Decimal | None
    passenger_name: str
    pax_type: PaxType
    staff_sign: str
    transaction_type: str
    booking_ref: str
    flags: tuple[str, ...]
    source_row: int
    tax: Decimal | None = None
    comm: Decimal | None = None
    form_of_payment: str | None = None
    confirmed: str | None = None


@dataclass(frozen=True)
class ParsedReport:
    header: ReportHeader
    tickets: tuple[ParsedTicket, ...]
    warnings: tuple[str, ...]


class _Warnings:
    """Keeps the first ``MAX_WARNINGS`` messages and counts the rest."""

    def __init__(self) -> None:
        self._items: list[str] = []
        self._dropped = 0

    def add(self, message: str) -> None:
        if len(self._items) < MAX_WARNINGS:
            self._items.append(message)
        else:
            self._dropped += 1

    def result(self) -> tuple[str, ...]:
        if self._dropped:
            return (*self._items, f"... and {self._dropped} more warnings")
        return tuple(self._items)


def _short(value: object) -> str:
    """repr of a cell value, cut to ``WARNING_VALUE_CHARS`` so warnings stay small."""
    text = repr(value)
    if len(text) <= WARNING_VALUE_CHARS:
        return text
    return text[: WARNING_VALUE_CHARS - 3] + "..."


def parse_report(source: Path | bytes | bytearray | BinaryIO) -> ParsedReport:
    data = _read_bytes(source)
    with _open_sheet(data) as ws:
        rows = _iter_rows(ws)
        head = list(itertools.islice(rows, HEADER_SCAN_ROWS))
        header_row, columns, unknown, truncated = _locate_columns(head)
        header = _parse_header(head[: header_row - 1])

        warnings = _Warnings()
        if truncated:
            warnings.add(
                f"only the first {MAX_COLUMNS} columns were read; "
                f"columns after the {MAX_COLUMNS}th were ignored"
            )
        if unknown:
            warnings.add(_unknown_columns_message(unknown))
        # Every known column has a key; a column the file lacks is read as an empty cell.
        all_keys = (*_REQUIRED.values(), *_OPTIONAL.values())
        tickets: list[ParsedTicket] = []
        index_by_no: dict[str, int] = {}
        non_blank = 0
        blank_row: tuple[Any, ...] | None = None
        # Rows after the header that were already pulled while scanning for it come first.
        data_rows = itertools.chain(head[header_row:], rows)
        for row_no, row in enumerate(data_rows, start=header_row + 1):
            if row is blank_row:  # openpyxl yields one shared tuple for every missing row
                continue
            cells = {
                key: (row[columns[key]] if key in columns and columns[key] < len(row) else None)
                for key in all_keys
            }
            if all(_is_blank(cells[key]) for key in _REQUIRED.values()):
                blank_row = row
                continue
            non_blank += 1
            if non_blank > MAX_DATA_ROWS:
                raise ParseError(f"too many rows (more than {MAX_DATA_ROWS} data rows)")
            ticket = _build_ticket(cells, row_no, warnings)
            if ticket is None:
                continue
            first = index_by_no.get(ticket.ticket_no)
            if first is None:
                index_by_no[ticket.ticket_no] = len(tickets)
                tickets.append(ticket)
                continue
            warnings.add(
                f"row {row_no}: duplicate DOC NUMBER {ticket.ticket_no} ignored (first one kept)"
            )
            kept = tickets[first]
            if FLAG_DUPLICATE_IN_FILE not in kept.flags:
                flags = tuple(sorted({*kept.flags, FLAG_DUPLICATE_IN_FILE}))
                tickets[first] = dataclasses.replace(kept, flags=flags)
    return ParsedReport(header=header, tickets=tuple(tickets), warnings=warnings.result())


# --- loading ------------------------------------------------------------------------


def _read_bytes(source: Path | bytes | bytearray | BinaryIO) -> bytes:
    too_big = ParseError("file is larger than the 10 MB limit")
    if isinstance(source, Path):
        try:
            if source.stat().st_size > MAX_FILE_BYTES:
                raise too_big
            return source.read_bytes()
        except FileNotFoundError as exc:
            raise ParseError(f"file not found: {source}") from exc
        except OSError as exc:
            raise ParseError(f"cannot read file: {exc}") from exc
    if isinstance(source, bytes | bytearray):
        data = bytes(source)
    else:
        data = source.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise too_big
    return data


def _check_archive(data: bytes) -> None:
    """Reject zip bombs from the central directory, before anything is decompressed."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
    except zipfile.BadZipFile as exc:
        raise ParseError("not a valid .xlsx workbook") from exc
    if len(infos) > MAX_ZIP_ENTRIES:
        raise ParseError(f"workbook has too many entries (limit {MAX_ZIP_ENTRIES})")
    sizes = [i.file_size for i in infos]
    if any(size > MAX_ENTRY_BYTES for size in sizes) or sum(sizes) > MAX_UNCOMPRESSED_BYTES:
        raise ParseError("file expands to an unreasonable size")


@contextmanager
def _open_sheet(data: bytes) -> Iterator[ReadOnlyWorksheet]:
    if not openpyxl.xml.DEFUSEDXML:
        raise RuntimeError("defusedxml is required (is OPENPYXL_DEFUSEDXML disabled?)")
    _check_archive(data)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ParseError("not a valid .xlsx workbook") from exc
    try:
        sheets = wb.worksheets
        ws = sheets[0] if sheets else None
        if not isinstance(ws, ReadOnlyWorksheet):
            raise ParseError("workbook has no worksheet")
        # Declared dimensions are untrusted: too small drops rows, too big wastes time.
        ws.reset_dimensions()
        yield ws
    finally:
        wb.close()


def _iter_rows(ws: ReadOnlyWorksheet) -> Iterator[tuple[Any, ...]]:
    try:
        # types-openpyxl declares iter_rows with a Worksheet-only self type; it is a stub bug.
        yield from ws.iter_rows(  # type: ignore[misc]
            min_row=1, max_col=MAX_COLUMNS, values_only=True
        )
    except Exception as exc:
        raise ParseError("worksheet is malformed or uses forbidden XML") from exc


# --- header block -------------------------------------------------------------------


def _norm(value: object) -> str:
    return " ".join(str(value).upper().split()) if value is not None else ""


def _locate_columns(
    rows: list[tuple[Any, ...]],
) -> tuple[int, dict[str, int], list[str], bool]:
    """The header row number, the column index of every known column, the header titles that
    are neither known nor deliberately ignored (cleaned, shortened, without duplicates), and
    whether the header fills every column that is read (so more may follow, unread)."""
    for number, row in enumerate(rows, start=1):
        if any(_norm(v) == "DOC NUMBER" for v in row):
            columns: dict[str, int] = {}
            unknown: list[str] = []
            for idx, value in enumerate(row):
                title = _norm(value)
                key = _REQUIRED.get(title) or _OPTIONAL.get(title)
                if key is not None:
                    columns.setdefault(key, idx)
                    continue
                shown = clean_text(title)[:UNKNOWN_COLUMN_CHARS].strip()
                if shown and shown not in _KNOWN_IGNORED and shown not in unknown:
                    unknown.append(shown)
            missing = [h for h, k in _REQUIRED.items() if k not in columns]
            if missing:
                raise ParseError(f"missing required column(s): {', '.join(missing)}")
            full = len(row) >= MAX_COLUMNS and not _is_blank(row[MAX_COLUMNS - 1])
            return number, columns, unknown, full
    raise ParseError(
        f"header row not found (no 'DOC NUMBER' column in the first {HEADER_SCAN_ROWS} rows)"
    )


def _unknown_columns_message(unknown: list[str]) -> str:
    shown = ", ".join(unknown[:MAX_UNKNOWN_COLUMNS])
    extra = len(unknown) - MAX_UNKNOWN_COLUMNS
    if extra > 0:
        shown += f" (+{extra} more)"
    return f"unknown columns ignored: {shown}"


def _label_value(rows: list[tuple[Any, ...]], label: str) -> Any:
    for row in rows:
        for idx, value in enumerate(row):
            if _norm(value) == label.upper() and idx + 1 < len(row):
                return row[idx + 1]
    return None


def _parse_header(rows: list[tuple[Any, ...]]) -> ReportHeader:
    range_value = _label_value(rows, "Report date range")
    current_value = _label_value(rows, "Current date")

    raw = str(range_value).strip() if range_value is not None else ""
    m = _RANGE_RE.match(raw)
    if not m or m.group(2).upper() not in _MONTHS:
        raise ParseError(f"'Report date range' must look like 01OCT, got {_short(raw)}")
    day, month = int(m.group(1)), _MONTHS[m.group(2).upper()]

    generated = _parse_current(current_value)
    year = generated.year - 1 if month > generated.month else generated.year
    try:
        report_date = date(year, month, day)
    except ValueError as exc:
        raise ParseError(
            f"'Report date range' {_short(raw)} is not a valid date in {year}"
        ) from exc
    return ReportHeader(report_range_raw=raw, report_date=report_date, generated_on=generated)


def _parse_current(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = str(value).strip() if value is not None else ""
    m = _CURRENT_RE.match(raw)
    if m and m.group(2).upper() in _MONTHS:
        try:
            return date(int(m.group(3)), _MONTHS[m.group(2).upper()], int(m.group(1)))
        except ValueError:
            pass
    raise ParseError(f"'Current date' must look like 02-OCT-2026, got {_short(raw)}")


# --- rows ---------------------------------------------------------------------------


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


_LINE_BREAKS = frozenset("\t\n\r\v\f\x1c\x1d\x1e\x85")


def clean_text(value: str) -> str:
    """``value`` without control, format (bidi, zero-width), surrogate, private-use or
    unassigned characters; line breaks and the line/paragraph separators become spaces.

    Cell text ends up in HTML and in logs: an override such as U+202E must not reorder what a
    manager reads. Ordinary letters (any script), digits, punctuation and spaces are kept.
    """
    if value.isprintable():
        return value.strip()
    out: list[str] = []
    for ch in value:
        category = unicodedata.category(ch)
        if ch in _LINE_BREAKS or category in ("Zl", "Zp"):
            out.append(" ")
        elif category[0] != "C":
            out.append(ch)
    return "".join(out).strip()


def _text(value: object, field: str, row_no: int, warnings: _Warnings) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return clean_text(value)
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    else:
        text = clean_text(str(value))
    warnings.add(f"row {row_no}: {field} was a number, converted to text {_short(text)}")
    return text


def _cost(value: object) -> tuple[Decimal | None, bool]:
    """Return ``(cost, was_rounded)``; cost is None when absent or not storable."""
    if value is None or isinstance(value, bool):
        return None, False
    try:
        if isinstance(value, int):
            parsed = Decimal(value)
        elif isinstance(value, float):
            parsed = Decimal(repr(value))
        elif isinstance(value, str):
            parsed = Decimal(value.replace(",", "").replace(" ", ""))
        else:
            return None, False
        if not parsed.is_finite() or abs(parsed) >= COST_LIMIT:
            return None, False
        rounded = parsed.quantize(_CENT, rounding=ROUND_HALF_UP)
    except ArithmeticError:
        return None, False
    return rounded, rounded != parsed


def _limited(
    text: str, key: str, label: str, row_no: int, flags: set[str], warnings: _Warnings
) -> str:
    limit = FIELD_LIMITS[key]
    if len(text) <= limit:
        return text
    flags.add(FLAG_FIELD_TOO_LONG)
    warnings.add(f"row {row_no}: {label} longer than {limit} characters, truncated")
    return text[:limit]


_PLAIN_AMOUNT_RE = re.compile(r"-?[0-9]+(?:\.[0-9]+)?", re.ASCII)
_COMMA_DECIMAL_RE = re.compile(r"(-?[0-9]+),([0-9]{1,2})", re.ASCII)
_COMMA_THOUSANDS_RE = re.compile(r"-?[0-9]{1,3}(?:,[0-9]{3})+(?:\.[0-9]+)?", re.ASCII)


def _amount_text(raw: str) -> str | None:
    """``raw`` as a plain decimal string, or None when it is ambiguous or not an amount.

    A trailing comma with one or two digits is the decimal separator ("350,5", "1 234,56"); a
    comma followed by exactly three digits is a thousands separator ("1,234", "1,234.56"). Spaces
    (and no-break spaces) are ignored. Anything else ("1.234,56", "1,2,3") is refused."""
    text = "".join(raw.split())
    if _PLAIN_AMOUNT_RE.fullmatch(text):
        return text
    decimal = _COMMA_DECIMAL_RE.fullmatch(text)
    if decimal:
        return f"{decimal.group(1)}.{decimal.group(2)}"
    if _COMMA_THOUSANDS_RE.fullmatch(text):
        return text.replace(",", "")
    return None


def _optional_money(value: object, label: str, row_no: int, warnings: _Warnings) -> Decimal | None:
    """An optional amount (TAX, COMM): None when the cell is empty. A cell that is not an
    amount, or is ambiguous, is ignored with a warning (it never fails the row); an amount with
    more than two decimals is rounded, with a warning. Unlike TOTAL DOC, a comma before one or
    two digits is a decimal comma."""
    if _is_blank(value):
        return None
    parsable: object = value
    if isinstance(value, str):
        parsable = _amount_text(value)
    amount, rounded = _cost(parsable)
    if amount is None:
        warnings.add(f"row {row_no}: {label} is not a valid amount, ignored")
    elif rounded:
        warnings.add(f"row {row_no}: {label} rounded to 2 decimals")
    return amount


def _build_ticket(cells: dict[str, Any], row_no: int, warnings: _Warnings) -> ParsedTicket | None:
    ticket_no = _text(cells["ticket_no"], "DOC NUMBER", row_no, warnings)
    if not ticket_no:
        warnings.add(f"row {row_no}: blank DOC NUMBER, row skipped")
        return None
    if not _TICKET_NO_RE.fullmatch(ticket_no):
        warnings.add(
            f"row {row_no}: DOC NUMBER {_short(ticket_no)} is not 8-15 digits, row skipped"
        )
        return None

    flags: set[str] = set()
    name = _text(cells["passenger_name"], "PAX NAME", row_no, warnings)
    pax_type: PaxType = "ADT"
    suffix = _PAX_SUFFIX_RE.search(name)
    if suffix:
        pax_type = "INF" if suffix.group(1).upper() == "INF" else "CHD"
        name = name[: suffix.start()].rstrip()
    elif _PAX_MARKER_RE.search(name):
        warnings.add(
            f"row {row_no}: pax type marker in {_short(name)} not recognised, treated as adult"
        )

    if len(name) > MAX_NAME_CHARS:
        flags.add(FLAG_FIELD_TOO_LONG)
        warnings.add(f"row {row_no}: PAX NAME longer than {MAX_NAME_CHARS} characters, truncated")
        name = name[:MAX_NAME_CHARS].rstrip()

    cost, rounded = _cost(cells["cost"])
    if cost is None:
        flags.add(FLAG_MISSING_COST)
    if rounded:
        flags.add(FLAG_COST_ROUNDED)

    def field(key: str, label: str) -> str:
        return _limited(
            _text(cells[key], label, row_no, warnings), key, label, row_no, flags, warnings
        )

    seq_no = field("seq_no", "SEQ NO")
    airline = field("airline", "A/L")
    staff_sign = field("staff_sign", "AS")
    trnc = field("transaction_type", "TRNC").upper()
    booking_ref = field("booking_ref", "RLOC")
    if trnc == TRNC_CANCELLED:
        flags.add(FLAG_CANCELLED)
    elif trnc != TRNC_ISSUED:
        flags.add(FLAG_UNKNOWN_TRNC)
    if not booking_ref:
        flags.add(FLAG_MISSING_BOOKING_REF)
    if not staff_sign:
        flags.add(FLAG_MISSING_STAFF_SIGN)

    tax = _optional_money(cells["tax"], "TAX", row_no, warnings)
    comm = _optional_money(cells["comm"], "COMM", row_no, warnings)
    fop = field("form_of_payment", "FP")
    confirmed = field("confirmed", "CONFIRMED")

    return ParsedTicket(
        seq_no=seq_no,
        airline=airline,
        ticket_no=ticket_no,
        cost=cost,
        passenger_name=name,
        pax_type=pax_type,
        staff_sign=staff_sign,
        transaction_type=trnc,
        booking_ref=booking_ref,
        flags=tuple(sorted(flags)),
        source_row=row_no,
        tax=tax,
        comm=comm,
        form_of_payment=fop or None,
        confirmed=confirmed or None,
    )

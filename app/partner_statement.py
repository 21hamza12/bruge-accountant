"""The partner statement (« relevé de compte partenaire ») as an Excel workbook (phase F, slice F3).

Plan doc 16 §3 and doc 06 §2: the agency header, the partner and the period, the tickets (date,
ticket number, airline, passenger, PNR, cost, fee, P.T), the manual lines (kind, reason, whether
they count), the payments (date, method, reference, amount), then the previous balance, the
totals and the new balance with its sign explained. Excel only (owner's answer 4).

**Where the figures come from.**

* A **closed** period is built ONLY from its frozen row: the lines from ``row.snapshot`` and the
  totals from the columns written by the close (``from_snapshot``). Nothing live is read for it,
  so a sale corrected after the close never changes the statement already sent (the difference
  is a ``CORRECTION`` line of the open period). The snapshot is checked to add up to the frozen
  totals; a snapshot that does not is refused (``StatementError``), never printed.
* An **open** period may be exported too, from the live account (``_from_live``), marked
  « PROVISOIRE » in its title, its status line and its file name.

**Formula injection.** Every text that a person typed (the partner's name, the passenger, the
reasons, the references, the agency's details) goes through ``export_apres.neutralise`` (a
leading ``= + - @``, tab or CR gets an apostrophe; characters XML cannot hold are dropped) and is
forced to a string cell. Amounts and dates are numbers and dates, never text; nothing is a
formula.

**Formatting.** Amounts are ``Decimal`` from end to end and written as numbers with two decimals
(``#,##0.00``), as ``export_apres`` does; dates are real dates shown ``dd/mm/yyyy``; the ticket
header row is frozen; the columns have fixed widths.

Who: manager and accountant (``permissions.can_export_partner_statement``); the agency is always
the actor's own (a partner of another agency is not found). This module only reads.
"""

from __future__ import annotations

import io
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from typing import Any, Final

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import agency_details, partner_book, permissions
from app.airlines import airline_name
from app.errors import UserError
from app.export_apres import MONEY_FORMAT, XLSX_MEDIA_TYPE, neutralise
from app.models import PERIOD_CLOSED, Agency, PartnerPeriod, User
from app.partner_book import KIND_LABELS, METHOD_LABELS, PeriodView
from app.partner_ledger import CENT, ZERO

TITLE: Final = "RELEVÉ DE COMPTE PARTENAIRE"
PROVISIONAL: Final = "PROVISOIRE"
PROVISIONAL_STATE: Final = "PROVISOIRE : période ouverte, les montants peuvent encore changer"
SIGN_RULE: Final = (
    "Solde positif : le partenaire doit à l'agence ; solde négatif : l'agence doit au partenaire."
)
OWES_AGENCY: Final = "Le partenaire doit ce montant à l'agence."
AGENCY_OWES: Final = "L'agence doit ce montant au partenaire."
SETTLED: Final = "Compte soldé."
COUNTED: Final = "comptée"
REJECTED: Final = "rejetée (non comptée)"
PENDING: Final = "en attente (non comptée)"
DATE_FORMAT: Final = "dd/mm/yyyy"
SHEET_TITLE: Final = "Relevé"
TICKET_HEAD: Final = ("Date", "N° billet", "Compagnie", "Passager", "PNR", "Coût", "Frais", "P.T")
WIDTHS: Final = (13, 30, 26, 30, 10, 15, 24, 16)
LAST_COL: Final = len(TICKET_HEAD)  # H: every amount of the lines, payments and summary

NO_PERMISSION = "Seuls le gérant et le comptable téléchargent les relevés des partenaires."
NOT_FOUND = "Période introuvable."
BAD_SNAPSHOT = "Le relevé figé de cette période est illisible ou ne correspond pas à ses totaux."

_FILL_HEAD = PatternFill("solid", fgColor="FFD9D9D9")
_BOLD = Font(bold=True)


class StatementError(UserError):
    """The statement cannot be built: a frozen period that does not add up (never printed)."""


class StatementNotFound(StatementError):
    """No such period for this partner (a 404)."""


class StatementForbidden(StatementError):
    """The actor may not download a statement (a 403)."""


@dataclass(frozen=True)
class TicketRow:
    day: date
    doc_number: str
    airline: str  # the airline's name (``app.airlines``)
    passenger: str
    pnr: str
    cost: Decimal
    fee: Decimal
    pt: Decimal  # cost + fee


@dataclass(frozen=True)
class EntryRow:
    day: date
    kind: str  # the kind's French label
    reason: str
    state: str  # COUNTED, REJECTED or PENDING
    amount: Decimal  # signed

    @property
    def counted(self) -> bool:
        return self.state == COUNTED


@dataclass(frozen=True)
class PaymentRow:
    day: date
    method: str  # the method's French label
    reference: str
    amount: Decimal


@dataclass(frozen=True)
class Body:
    """The lines and the totals of one period."""

    tickets: tuple[TicketRow, ...]
    entries: tuple[EntryRow, ...]
    payments: tuple[PaymentRow, ...]
    opening: Decimal
    sales_total: Decimal
    entries_total: Decimal
    payments_total: Decimal
    closing: Decimal


@dataclass(frozen=True)
class Statement:
    header: tuple[str, ...]  # the agency's lines, missing details left out
    currency: str
    partner_id: int
    partner_name: str
    start: date
    end: date
    provisional: bool  # an open period, from the live account
    closed_on: date | None
    issued_on: date
    period_id: int | None  # the frozen row of a closed period
    body: Body

    @property
    def closing(self) -> Decimal:
        return self.body.closing


# --- the figures ----------------------------------------------------------------------------------


def _amount(raw: object) -> Decimal:
    if not isinstance(raw, str | int | Decimal) or isinstance(raw, bool):
        raise StatementError(BAD_SNAPSHOT)
    try:
        value = Decimal(raw)
    except InvalidOperation:
        raise StatementError(BAD_SNAPSHOT) from None
    if not value.is_finite():
        raise StatementError(BAD_SNAPSHOT)
    return value.quantize(CENT)


def _day(raw: object) -> date:
    if not isinstance(raw, str):
        raise StatementError(BAD_SNAPSHOT)
    try:
        return date.fromisoformat(raw)
    except ValueError:
        raise StatementError(BAD_SNAPSHOT) from None


def _str(raw: object) -> str:
    return raw if isinstance(raw, str) else ""


def _list(raw: object) -> list[Mapping[str, Any]]:
    if not isinstance(raw, list) or not all(isinstance(item, Mapping) for item in raw):
        raise StatementError(BAD_SNAPSHOT)
    return raw


def _ticket(
    day: date, fee: Decimal, doc: str, airline: str, pax: str, pnr: str, cost: Decimal
) -> TicketRow:
    return TicketRow(day, doc, airline_name(airline), pax, pnr, cost, fee, cost + fee)


def _state(counted: bool, decision: str | None) -> str:
    if counted:
        return COUNTED
    return REJECTED if decision == "reject" else PENDING


def _checked(body: Body) -> Body:
    """The lines add up to the totals and the totals to the new balance; else StatementError."""
    if (
        sum((t.pt for t in body.tickets), ZERO) != body.sales_total
        or sum((e.amount for e in body.entries if e.counted), ZERO) != body.entries_total
        or sum((p.amount for p in body.payments), ZERO) != body.payments_total
        or body.opening + body.sales_total + body.entries_total - body.payments_total
        != body.closing
    ):
        raise StatementError(BAD_SNAPSHOT)
    return body


def from_snapshot(row: PartnerPeriod, snapshot: Mapping[str, Any] | None) -> Body:
    """The body of a CLOSED period from its frozen row only: the totals written by the close and
    the lines of its snapshot (which must agree with them)."""
    frozen = (
        row.opening_balance,
        row.sales_total,
        row.entries_total,
        row.payments_total,
        row.closing_balance,
    )
    if row.status != PERIOD_CLOSED or snapshot is None or any(v is None for v in frozen):
        raise StatementError(BAD_SNAPSHOT)
    opening, sales_total, entries_total, payments_total, closing = (
        Decimal(v).quantize(CENT) for v in frozen if v is not None
    )
    try:
        stored = tuple(
            _amount(snapshot[k])
            for k in ("opening", "sales_total", "entries_total", "payments_total", "closing")
        )
        tickets = tuple(
            _ticket(
                _day(sale["day"]),
                _amount(sale["fee"]),
                _str(t.get("doc_number")),
                _str(t.get("airline")),
                _str(t.get("pax_name")),
                _str(t.get("rloc")),
                _amount(t["total_doc"]),
            )
            for sale in _list(snapshot["sales"])
            for t in _list(sale["tickets"])
        )
        entries = tuple(
            EntryRow(
                _day(e["date"]),
                KIND_LABELS.get(_str(e.get("kind")), _str(e.get("kind"))),
                _str(e.get("reason")),
                _state(e.get("counted") is True, e.get("decision")),
                _amount(e["amount"]),
            )
            for e in _list(snapshot["entries"])
        )
        payments = tuple(
            PaymentRow(
                _day(p["date"]),
                METHOD_LABELS.get(_str(p.get("method")), _str(p.get("method"))),
                _str(p.get("reference")),
                _amount(p["amount"]),
            )
            for p in _list(snapshot["payments"])
        )
    except (KeyError, TypeError, AttributeError):
        raise StatementError(BAD_SNAPSHOT) from None
    if stored != (opening, sales_total, entries_total, payments_total, closing):
        raise StatementError(BAD_SNAPSHOT)
    return _checked(
        Body(
            tickets,
            entries,
            payments,
            opening,
            sales_total,
            entries_total,
            payments_total,
            closing,
        )  # fmt: skip
    )


def _from_live(period: PeriodView) -> Body:
    """The body of an OPEN period, as the account computes it now."""
    totals = period.totals
    if totals is None:  # pragma: no cover - an open period always has its live totals
        raise StatementError(BAD_SNAPSHOT)
    tickets = tuple(
        _ticket(line.day, line.fee, t.doc_number, t.airline, t.pax_name, t.rloc, t.total_doc)
        for line in period.sales
        for t in line.tickets
    )
    entries = tuple(
        EntryRow(
            e.entry.entry_date,
            KIND_LABELS.get(e.entry.kind, e.entry.kind),
            e.entry.reason or "",
            _state(e.counted, None if e.decision is None else e.decision.decision),
            e.entry.amount,
        )
        for e in period.entries
    )
    payments = tuple(
        PaymentRow(p.paid_on, METHOD_LABELS.get(p.method, p.method), p.reference or "", p.amount)
        for p in period.payments
    )
    return _checked(
        Body(
            tickets,
            entries,
            payments,
            totals.opening,
            totals.sales_total,
            totals.entries_total,
            totals.payments_total,
            totals.closing,
        )  # fmt: skip
    )


def statement_for(
    db: Session, actor: User, partner_id: int, start: date, *, today: date
) -> Statement:
    """The statement of the period of ``partner_id`` (of the actor's agency) starting on
    ``start``. ``partner_book.AccountNotFound`` for a partner that is not the agency's,
    ``StatementNotFound`` for a period that does not exist."""
    if not permissions.can_export_partner_statement(actor):
        raise StatementForbidden(NO_PERMISSION)
    acc = partner_book.get_account(db, actor, partner_id, today, snapshots=True)
    period = next((p for p in acc.periods if p.start == start), None)
    if period is None:
        raise StatementNotFound(NOT_FOUND)
    if period.closed and period.row is not None:
        body = from_snapshot(period.row, period.row.snapshot)
        closed_at = period.row.closed_at
        closed_on = closed_at.date() if closed_at is not None else None
    else:
        body, closed_on = _from_live(period), None
    agency = db.scalars(select(Agency).where(Agency.id == actor.agency_id)).one()
    details = agency_details.get_details(db, actor.agency_id)
    return Statement(
        header=tuple(agency_details.header_lines(details, agency.name)),
        currency=agency.currency,
        partner_id=acc.partner.id,
        partner_name=acc.partner.name,
        start=period.start,
        end=period.end,
        provisional=not period.closed,
        closed_on=closed_on,
        issued_on=today,
        period_id=period.row.id if period.closed and period.row is not None else None,
        body=body,
    )


def statement_filename(statement: Statement) -> str:
    """``releve-partenaire-<id>-<start>-<end>.xlsx`` (``-PROVISOIRE`` for an open period): ids
    and dates only, never a name."""
    suffix = f"-{PROVISIONAL}" if statement.provisional else ""
    return (
        f"releve-partenaire-{statement.partner_id}-{statement.start.isoformat()}-"
        f"{statement.end.isoformat()}{suffix}.xlsx"
    )


# --- the workbook ---------------------------------------------------------------------------------


def _fr(day: date) -> str:
    return day.strftime("%d/%m/%Y")


class _Sheet:
    """A cursor over the sheet: one ``row`` call per line, top to bottom."""

    def __init__(self, ws: Worksheet) -> None:
        self.ws = ws
        self.at = 0

    def skip(self) -> None:
        self.at += 1

    def text(self, col: int, value: str, *, bold: bool = False) -> None:
        clean = neutralise(value)
        if not clean:
            return
        cell = self.ws.cell(row=self.at, column=col)
        cell.value = clean
        cell.data_type = "s"  # a string cell, whatever the text looks like
        if bold:
            cell.font = _BOLD

    def money(self, col: int, value: Decimal, *, bold: bool = False) -> None:
        cell = self.ws.cell(row=self.at, column=col)
        cell.value = value.quantize(CENT)
        cell.number_format = MONEY_FORMAT
        if bold:
            cell.font = _BOLD

    def day(self, col: int, value: date) -> None:
        cell = self.ws.cell(row=self.at, column=col)
        cell.value = value
        cell.number_format = DATE_FORMAT

    def line(self, *texts: str, bold: bool = False) -> None:
        self.at += 1
        for col, value in enumerate(texts, start=1):
            self.text(col, value, bold=bold and col == 1)

    def head(self, titles: Sequence[str | None]) -> None:
        self.at += 1
        for col, title in enumerate(titles, start=1):
            if title is None:
                continue
            cell = self.ws.cell(row=self.at, column=col, value=title)
            cell.font, cell.fill = _BOLD, _FILL_HEAD
            cell.alignment = Alignment(horizontal="center")

    def total(self, label: str, value: Decimal, *, col: int = LAST_COL) -> None:
        self.at += 1
        self.text(1, label, bold=True)
        self.money(col, value, bold=True)


def build_statement_xlsx(statement: Statement) -> bytes:
    """The workbook of ``statement`` as bytes (pure: no database)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None  # noqa: S101 - a new workbook always has an active sheet
    ws.title = SHEET_TITLE
    sheet = _Sheet(ws)
    body = statement.body
    # The agency.
    for n, line in enumerate(statement.header):
        sheet.line(line, bold=n == 0)
    if statement.header:
        ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    sheet.skip()
    # The partner and the period.
    title = f"{TITLE} — {PROVISIONAL}" if statement.provisional else TITLE
    sheet.line(title, bold=True)
    ws.cell(row=sheet.at, column=1).font = Font(bold=True, size=13)
    sheet.line("Partenaire", statement.partner_name, bold=True)
    sheet.line("Période", f"du {_fr(statement.start)} au {_fr(statement.end)}", bold=True)
    if statement.provisional or statement.closed_on is None:
        state = PROVISIONAL_STATE
    else:
        state = f"Clôturée le {_fr(statement.closed_on)}"
    sheet.line("État", state, bold=True)
    sheet.line("Édité le", _fr(statement.issued_on), bold=True)
    sheet.line("Devise", statement.currency, bold=True)
    sheet.skip()
    # The tickets.
    sheet.head(TICKET_HEAD)
    ws.freeze_panes = f"A{sheet.at + 1}"
    if not body.tickets:
        sheet.line("Aucun billet dans cette période.")
    for t in body.tickets:
        sheet.at += 1
        sheet.day(1, t.day)
        sheet.text(2, t.doc_number)
        sheet.text(3, t.airline)
        sheet.text(4, t.passenger)
        sheet.text(5, t.pnr)
        sheet.money(6, t.cost)
        sheet.money(7, t.fee)
        sheet.money(8, t.pt)
    sheet.at += 1
    sheet.text(1, "Total des billets", bold=True)
    sheet.money(6, sum((t.cost for t in body.tickets), ZERO), bold=True)
    sheet.money(7, sum((t.fee for t in body.tickets), ZERO), bold=True)
    sheet.money(8, body.sales_total, bold=True)
    sheet.skip()
    # The manual lines.
    sheet.line("Lignes manuelles", bold=True)
    sheet.head(("Date", "Type", "Motif", None, None, None, "État", "Montant"))
    if not body.entries:
        sheet.line("Aucune ligne.")
    for e in body.entries:
        sheet.at += 1
        sheet.day(1, e.day)
        sheet.text(2, e.kind)
        sheet.text(3, e.reason)
        sheet.text(7, e.state)
        sheet.money(8, e.amount)
    sheet.total("Total des lignes comptées", body.entries_total)
    sheet.skip()
    # The payments.
    sheet.line("Paiements", bold=True)
    sheet.head(("Date", "Mode", "Référence", None, None, None, None, "Montant"))
    if not body.payments:
        sheet.line("Aucun paiement.")
    for p in body.payments:
        sheet.at += 1
        sheet.day(1, p.day)
        sheet.text(2, p.method)
        sheet.text(3, p.reference)
        sheet.money(8, p.amount)
    sheet.total("Total des paiements", body.payments_total)
    sheet.skip()
    # The balance.
    sheet.line("Récapitulatif", bold=True)
    sheet.total("Solde précédent", body.opening)
    sheet.total("Billets (+)", body.sales_total)
    sheet.total("Lignes manuelles (+/−)", body.entries_total)
    sheet.total("Paiements (−)", body.payments_total)
    sheet.total("Nouveau solde", body.closing)
    if body.closing > 0:
        meaning = OWES_AGENCY
    elif body.closing < 0:
        meaning = AGENCY_OWES
    else:
        meaning = SETTLED
    sheet.line("Sens du solde", meaning, bold=True)
    sheet.line(SIGN_RULE)
    for n, width in enumerate(WIDTHS):
        ws.column_dimensions[chr(ord("A") + n)].width = width
    ws.page_setup.orientation = "landscape"
    fit = ws.sheet_properties.pageSetUpPr
    if fit is not None:  # printed one page wide, as many pages long as needed
        fit.fitToPage = True
        ws.page_setup.fitToHeight = 0
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


__all__ = [
    "PROVISIONAL",
    "TITLE",
    "XLSX_MEDIA_TYPE",
    "Body",
    "EntryRow",
    "PaymentRow",
    "Statement",
    "StatementError",
    "StatementForbidden",
    "StatementNotFound",
    "TicketRow",
    "build_statement_xlsx",
    "from_snapshot",
    "statement_filename",
    "statement_for",
]

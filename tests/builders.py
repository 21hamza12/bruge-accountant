"""Helpers that build small Amadeus-shaped workbooks for tests."""

from __future__ import annotations

import io
import re
import zipfile
from collections.abc import Callable, Mapping, Sequence
from typing import Any

import openpyxl

AMADEUS_HEADERS = ["SEQ NO", "A/L", "DOC NUMBER", "TOTAL DOC", "PAX NAME", "AS", "TRNC", "RLOC"]


def ticket_row(
    seq: Any = "000001",
    airline: Any = "199",
    doc: Any = "3940000001",
    cost: Any = 1000,
    name: Any = "DOE/JOHN",
    staff: Any = "0004SV",
    trnc: Any = "TKTT",
    rloc: Any = "AAAAAA",
) -> list[Any]:
    return [seq, airline, doc, cost, name, staff, trnc, rloc]


def build_workbook(
    rows: list[list[Any]],
    *,
    range_raw: str | None = "01OCT",
    current: Any = "02-OCT-2026",
    headers: list[str] | None = None,
    header_row: int = 5,
    hand_typed: bool = False,
    stray: dict[str, Any] | None = None,
) -> bytes:
    """Create an .xlsx (as bytes) shaped like an Amadeus daily sales export."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = "Agy no"
    ws["B1"] = ""
    ws["C1"] = "Report date range"
    if range_raw is not None:
        ws["D1"] = range_raw
    ws["A2"] = "Office"
    ws["A3"] = "Agent"
    ws["C3"] = "Current date"
    if current is not None:
        ws["D3"] = current
    for col, title in enumerate(AMADEUS_HEADERS if headers is None else headers, start=1):
        ws.cell(row=header_row, column=col, value=title)
    if hand_typed:
        for col, title in enumerate(["P.B", "AG", "CLT", "AP"], start=9):
            ws.cell(row=header_row, column=col, value=title)
    for offset, values in enumerate(rows, start=1):
        for col, value in enumerate(values, start=1):
            if value is not None:
                ws.cell(row=header_row + offset, column=col, value=value)
        if hand_typed:
            ws.cell(row=header_row + offset, column=9, value=999)
            ws.cell(row=header_row + offset, column=10, value="JUNK")
    for coordinate, value in (stray or {}).items():
        ws[coordinate] = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def set_cached_formula(xlsx: bytes, cell: str, formula: str, cached: int | None) -> bytes:
    """Rewrite sheet XML so ``cell`` holds ``<f>formula</f>`` plus an optional cached ``<v>``.

    openpyxl cannot write cached formula results, so we patch the XML directly.
    """
    src = zipfile.ZipFile(io.BytesIO(xlsx))
    out = io.BytesIO()
    pattern = re.compile(rf'<c r="{cell}"[^>]*?(/>|>.*?</c>)', re.DOTALL)
    value_xml = "" if cached is None else f"<v>{cached}</v>"
    replacement = f'<c r="{cell}"><f>{formula}</f>{value_xml}</c>'
    patched = False
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                text, count = pattern.subn(replacement, data.decode("utf-8"), count=1)
                patched = count == 1
                data = text.encode("utf-8")
            dst.writestr(item, data)
    assert patched, f"cell {cell} not found in sheet XML"
    return out.getvalue()


def patch_sheet_xml(xlsx: bytes, transform: Callable[[str], str]) -> bytes:
    """Return a copy of ``xlsx`` whose first sheet XML went through ``transform``."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(xlsx)) as src, zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = transform(data.decode("utf-8")).encode("utf-8")
            dst.writestr(item, data)
    return out.getvalue()


def add_zip_entries(xlsx: bytes, entries: dict[str, bytes]) -> bytes:
    """Return a copy of ``xlsx`` with extra (compressed) entries."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(xlsx)) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        for name, data in entries.items():
            dst.writestr(name, data)
    return out.getvalue()


# --- the newer 16-column Amadeus export -----------------------------------------------------

AMADEUS_16_HEADERS = [
    "SEQ NO",
    "CONFIRMED",
    "A/L",
    "DOC NUMBER",
    "TOTAL DOC",
    "TAX",
    "FEE",
    "COMM",
    "AGENT",
    "FP",
    "PAX NAME",
    "AS",
    "RLOC",
    "TRNC",
    "REPORT IND",
    "NDC",
]


def amadeus_line(
    seq: Any = "000001",
    airline: Any = "199",
    doc: Any = "1990000001",
    total: Any = 1000,
    name: Any = "SYNTH01/PAX",
    staff: Any = "0004SV",
    trnc: Any = "TKTT",
    rloc: Any = "AAAAAA",
    **others: Any,
) -> dict[str, Any]:
    """One line keyed by header title; TAX, FEE, COMM, FP, CONFIRMED... go in ``others``
    with the title as key (``**{"FP": "CA"}``)."""
    line: dict[str, Any] = {
        "SEQ NO": seq,
        "A/L": airline,
        "DOC NUMBER": doc,
        "TOTAL DOC": total,
        "PAX NAME": name,
        "AS": staff,
        "TRNC": trnc,
        "RLOC": rloc,
    }
    line.update(others)
    return line


def build_amadeus_xlsx(
    lines: Sequence[Mapping[str, Any]],
    *,
    header: Sequence[str] | None = None,
    extra_columns: Mapping[str, Sequence[Any]] | None = None,
    range_raw: str | None = "20OCT",
    current: Any = "21-OCT-2026",
) -> bytes:
    """An .xlsx in the 16-column layout of the real export (or in the column order ``header``).

    ``lines`` are dicts keyed by header title (see ``amadeus_line``); a title missing from a line
    leaves the cell empty. ``extra_columns`` adds columns after the header: title -> one value per
    line. Rows 1 to 3 hold Agy no / Office / Agent, the report date range, the current date and
    the currency, like the real export.
    """
    titles = list(AMADEUS_16_HEADERS if header is None else header)
    extras = dict(extra_columns or {})
    rows = [
        [line.get(title) for title in titles] + [values[i] for values in extras.values()]
        for i, line in enumerate(lines)
    ]
    return build_workbook(
        rows,
        range_raw=range_raw,
        current=current,
        headers=[*titles, *extras],
        stray={"F1": "Currency", "G1": "MRU"},
    )

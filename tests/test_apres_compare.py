"""The parallel-run comparison of two APRES workbooks (``app.apres_compare``), on synthetic data.

The app's side is a workbook built by ``export_apres.build_apres_xlsx`` (what is stored at close);
her side is the same workbook after the edits a person could make. Both go through ``parse_apres``,
exactly as in the web route.
"""

from __future__ import annotations

import io
from collections.abc import Callable
from datetime import date
from decimal import Decimal
from typing import Any

import openpyxl
import pytest

from app import apres_compare, export_apres
from app.export_apres import ApresLine
from app.parsing.amadeus import parse_apres

DAY = date(2026, 10, 20)
FIRST = 6  # first data row of the workbook
COLS = {"SEQ NO": 1, "A/L": 2, "DOC NUMBER": 3, "TOTAL DOC": 4, "PAX NAME": 5, "AS": 6,
        "TRNC": 7, "RLOC": 8, "P.B": 9, "AG": 10, "CLT": 11, "AP": 12}  # fmt: skip


def line(n: int, **kw: Any) -> ApresLine:
    base: dict[str, Any] = {
        "seq_no": f"{n:06d}",
        "airline": "199",
        "doc_number": f"{3_940_000_000 + n}",
        "total_doc": Decimal(1000 + n),
        "pax_name": f"SYNTH{n:02d}/PAX",
        "staff_sign": "0004SV",
        "trnc": "TKTT",
        "rloc": f"SY{n:04d}",
        "pb": Decimal(1500 + n),
        "ag": "BRUGE",
        "clt": f"Client {n}",
        "ap": "+22212345000",
        "sale_id": n,
        "ticket_count": 1,
    }
    base.update(kw)
    return ApresLine(**base)


def app_book(count: int = 5) -> bytes:
    report = export_apres.make_report(
        day=DAY,
        agy_no="",
        range_raw="20OCT",
        current_date="21-OCT-2026",
        lines=[line(n) for n in range(1, count + 1)],
    )
    return export_apres.build_apres_xlsx(report)


def edited(data: bytes, change: Callable[[Any], None]) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws is not None
    change(ws)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def put(row: int, column: int, value: Any) -> Callable[[Any], None]:
    def change(ws: Any) -> None:
        ws.cell(row=row, column=column).value = value

    return change


def run(app: bytes, hers: bytes) -> apres_compare.Comparison:
    return apres_compare.compare(parse_apres(app), parse_apres(hers))


def test_identical_workbooks_have_no_difference() -> None:
    result = run(app_book(), app_book())
    assert (result.matched, result.identical, result.differing_cells) == (5, 5, 0)
    assert result.only_hers_count == result.only_app_count == 0
    assert result.same


def test_a_changed_amount_is_one_cell_with_both_values() -> None:
    hers = edited(app_book(), put(FIRST + 1, COLS["P.B"], 1700))
    result = run(app_book(), hers)
    assert (result.matched, result.identical, result.differing_cells) == (5, 4, 1)
    (diff,) = result.differences
    assert (diff.column, diff.hers, diff.app) == ("P.B", "1700.00", "1502.00")
    assert diff.doc_number == "3940000002"


def test_money_is_compared_as_a_decimal_not_as_text() -> None:
    def retype(ws: Any) -> None:
        ws.cell(row=FIRST, column=COLS["P.B"]).value = 1501  # the app's 1501.00
        ws.cell(row=FIRST + 1, column=COLS["P.B"]).value = "1 502,00"  # text, same amount

    assert run(app_book(), edited(app_book(), retype)).same


def test_a_cent_is_a_difference() -> None:
    hers = edited(app_book(), put(FIRST, COLS["TOTAL DOC"], 1001.01))
    (diff,) = run(app_book(), hers).differences
    assert (diff.column, diff.hers, diff.app) == ("TOTAL DOC", "1001.01", "1001.00")


def test_a_changed_name_is_reported_with_both_names() -> None:
    hers = edited(app_book(), put(FIRST + 2, COLS["PAX NAME"], "OTHER/NAME"))
    (diff,) = run(app_book(), hers).differences
    assert (diff.column, diff.hers, diff.app) == ("PAX NAME", "OTHER/NAME", "SYNTH03/PAX")


def test_text_is_compared_after_whitespace_and_case_normalisation() -> None:
    def retype(ws: Any) -> None:
        ws.cell(row=FIRST, column=COLS["PAX NAME"]).value = "  synth01 / pax  "
        ws.cell(row=FIRST, column=COLS["CLT"]).value = "CLIENT   1"
        ws.cell(row=FIRST + 1, column=COLS["AG"]).value = " bruge"

    assert run(app_book(), edited(app_book(), retype)).same


def test_the_child_mark_padding_is_not_a_difference() -> None:
    mine = export_apres.make_report(
        day=DAY, agy_no="", range_raw="20OCT",
        current_date="21-OCT-2026",
        lines=[line(1, pax_name="SYNTH/KID".ljust(24) + "(CHD)")],
    )  # fmt: skip
    hers = export_apres.make_report(
        day=mine.day, agy_no="", range_raw="20OCT", current_date="21-OCT-2026",
        lines=[line(1, pax_name="SYNTH/KID (CHD)")],
    )  # fmt: skip
    result = run(export_apres.build_apres_xlsx(mine), export_apres.build_apres_xlsx(hers))
    assert result.same


def test_a_child_in_hers_and_an_adult_in_the_app_differ() -> None:
    kid = export_apres.make_report(
        day=DAY, agy_no="", range_raw="20OCT",
        current_date="21-OCT-2026", lines=[line(1, pax_name="SYNTH/KID".ljust(24) + "(CHD)")],
    )  # fmt: skip
    adult = export_apres.make_report(
        day=kid.day, agy_no="", range_raw="20OCT", current_date="21-OCT-2026",
        lines=[line(1, pax_name="SYNTH/KID")],
    )  # fmt: skip
    (diff,) = run(
        export_apres.build_apres_xlsx(adult), export_apres.build_apres_xlsx(kid)
    ).differences
    assert diff.column == "PAX NAME" and "(CHD)" in diff.hers and "(CHD)" not in diff.app


def test_a_row_only_in_hers_and_a_row_only_in_the_app() -> None:
    def change(ws: Any) -> None:
        ws.delete_rows(FIRST + 1)  # the app's second row is missing from hers
        ws.cell(row=FIRST + 4, column=COLS["SEQ NO"]).value = "000099"  # a row only she has
        ws.cell(row=FIRST + 4, column=COLS["A/L"]).value = "199"
        ws.cell(row=FIRST + 4, column=COLS["DOC NUMBER"]).value = "3949999999"
        ws.cell(row=FIRST + 4, column=COLS["TOTAL DOC"]).value = 10
        ws.cell(row=FIRST + 4, column=COLS["PAX NAME"]).value = "ONLY/HERS"
        ws.cell(row=FIRST + 4, column=COLS["TRNC"]).value = "TKTT"

    result = run(app_book(), edited(app_book(), change))
    assert (result.matched, result.identical, result.differing_cells) == (4, 4, 0)
    assert [r.doc_number for r in result.only_app] == ["3940000002"]
    assert [r.doc_number for r in result.only_hers] == ["3949999999"]
    assert result.only_hers[0].pax_name == "ONLY/HERS"
    assert (result.only_hers_count, result.only_app_count) == (1, 1)
    assert not result.same


def test_the_thirteen_digit_ticket_form_aligns_with_the_ten_digit_one() -> None:
    hers = edited(app_book(), put(FIRST + 2, COLS["DOC NUMBER"], "1993940000003"))
    result = run(app_book(), hers)
    assert result.same and result.matched == 5


def test_an_unusable_ticket_number_does_not_align() -> None:
    hers = edited(app_book(), put(FIRST, COLS["DOC NUMBER"], "99940000001"))
    result = run(app_book(), hers)
    assert result.matched == 4 and result.only_hers_count == result.only_app_count == 1


def test_blank_against_zero_is_a_difference() -> None:
    hers = edited(app_book(), put(FIRST, COLS["P.B"], None))
    (diff,) = run(app_book(), hers).differences
    assert (diff.column, diff.hers, diff.app) == ("P.B", "", "1501.00")


def test_the_totals_row_and_the_header_block_are_not_rows() -> None:
    result = run(app_book(), app_book())
    assert result.app_rows == result.hers_rows == 5


def test_a_formula_cell_of_hers_reads_as_empty_and_the_export_guard_is_not_shown() -> None:
    mine = export_apres.make_report(
        day=DAY, agy_no="", range_raw="20OCT",
        current_date="21-OCT-2026", lines=[line(1, clt="=SUM(A1)")],
    )  # fmt: skip
    hers = edited(
        export_apres.build_apres_xlsx(mine),
        put(FIRST, COLS["CLT"], "=SUM(A1)"),
    )
    # her cell holds a formula with no stored result: it reads as empty, a visible difference
    (diff,) = run(export_apres.build_apres_xlsx(mine), hers).differences
    assert diff.column == "CLT" and diff.hers == "" and diff.app == "=SUM(A1)"


def test_the_kept_lists_are_bounded_but_the_counts_are_exact(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(apres_compare, "MAX_KEPT", 3)

    def change(ws: Any) -> None:
        for n in range(5):
            ws.cell(row=FIRST + n, column=COLS["P.B"]).value = 9
            ws.cell(row=FIRST + n, column=COLS["CLT"]).value = "x"

    result = run(app_book(), edited(app_book(), change))
    assert result.differing_cells == 10 and len(result.differences) == 3
    assert result.differences_cut


def test_cells_of_a_row_come_in_column_order() -> None:
    def change(ws: Any) -> None:
        ws.cell(row=FIRST, column=COLS["CLT"]).value = "x"
        ws.cell(row=FIRST, column=COLS["TOTAL DOC"]).value = 5
        ws.cell(row=FIRST, column=COLS["RLOC"]).value = "ZZZZZZ"

    result = run(app_book(), edited(app_book(), change))
    assert [d.column for d in result.differences] == ["TOTAL DOC", "RLOC", "CLT"]


def test_a_row_the_app_sold_nothing_for_compares_empty_cells() -> None:
    unmatched = export_apres.make_report(
        day=DAY, agy_no="", range_raw="20OCT",
        current_date="21-OCT-2026",
        lines=[line(1, pb=None, ag="", clt="", ap="", sale_id=None)],
    )  # fmt: skip
    book = export_apres.build_apres_xlsx(unmatched)
    assert run(book, book).same


def test_the_columns_compared_are_the_twelve_of_the_export() -> None:
    assert tuple(export_apres.HEADERS) == apres_compare.COLUMNS


def test_a_negative_zero_is_zero() -> None:
    hers = edited(app_book(), put(FIRST, COLS["P.B"], -0.004))
    base = edited(app_book(), put(FIRST, COLS["P.B"], 0))
    assert run(base, hers).same


def test_an_unreadable_amount_shows_her_own_text() -> None:
    def change(ws: Any) -> None:
        ws.cell(row=FIRST, column=COLS["P.B"]).value = "abc"
        ws.cell(row=FIRST + 1, column=COLS["TOTAL DOC"]).value = "12 EUR"

    diffs = run(app_book(), edited(app_book(), change)).differences
    assert [(d.column, d.hers) for d in diffs] == [("TOTAL DOC", "12 EUR"), ("P.B", "abc")] or [
        (d.column, d.hers) for d in diffs
    ] == [("P.B", "abc"), ("TOTAL DOC", "12 EUR")]


def test_a_truncated_cell_is_marked_and_never_equal() -> None:
    long = "C" * 250
    both = edited(app_book(), put(FIRST, COLS["CLT"], long))
    result = run(both, both)
    (diff,) = result.differences
    assert (
        diff.column == "CLT" and diff.hers.endswith("(tronqué)") and diff.app.endswith("(tronqué)")
    )
    only_hers = run(app_book(), both).differences
    assert only_hers[0].hers.endswith("… (tronqué)") and not only_hers[0].app.endswith("(tronqué)")

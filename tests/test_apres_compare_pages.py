"""The parallel-run page: ``/days/{day}/compare`` (C1 item 4), on synthetic data.

The accountant uploads her own APRES for a CLOSED day with a stored export; the page shows the
rows only in hers, only in the app, and each differing cell. Nothing is stored but an audit event
with counts. Most tests parse in-process (the same function the sandbox child runs); a few go
through the real sandbox.
"""

from __future__ import annotations

import io
import os
import tempfile
from collections.abc import Iterator
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import audit, permissions
from app.auth import ratelimit
from app.models import (
    ROLE_ACCOUNTANT,
    ROLE_MANAGER,
    ROLE_STAFF,
    AmadeusLine,
    DayExport,
    Upload,
    User,
)
from app.parsing.amadeus import MAX_FILE_BYTES, parse_apres
from app.parsing.sandbox import SandboxBusy
from app.web.routes import apres_compare as compare_routes
from tests.acceptance_support import text_of
from tests.builders import add_zip_entries
from tests.page_support import Env
from tests.test_export_apres import D20, close, fixture_tktt, sell_line, upload_20oct
from tests.test_query_counts import counting
from tests.test_sale_pages import SIGN, accountant, enter, table, tk
from tests.web_support import WebDb, audit_events, csrf_for

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FIRST = 6
COLS = {"SEQ NO": 1, "A/L": 2, "DOC NUMBER": 3, "TOTAL DOC": 4, "PAX NAME": 5, "AS": 6,
        "TRNC": 7, "RLOC": 8, "P.B": 9, "AG": 10, "CLT": 11, "AP": 12}  # fmt: skip
PATH = f"/days/{D20}/compare"


# --- helpers ------------------------------------------------------------------------------------


@pytest.fixture
def env(web_db: WebDb) -> Env:
    e = Env(web_db)
    e.factory.app.dependency_overrides[compare_routes.get_apres_parser] = (  # type: ignore[attr-defined]
        lambda: parse_apres
    )
    return e


@pytest.fixture
def closed(env: Env, gate_open: None) -> Iterator[tuple[Env, TestClient, bytes]]:
    """A closed day with its stored export; the client is the accountant who closed it."""
    upload_20oct(env)
    acc = accountant(env)
    sell_line(acc, 0)
    sell_line(acc, 1)
    assert close(acc).status_code == 303
    (stored,) = table(env, DayExport)
    yield env, acc, stored.content


def edited(data: bytes, *changes: tuple[int, str, Any]) -> bytes:
    """``data`` with cells set: ``(data row index, column title, value)``."""
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws is not None
    for index, title, value in changes:
        ws.cell(row=FIRST + index, column=COLS[title]).value = value
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def submit(
    c: TestClient, data: bytes | None, *, name: str = "sa feuille.xlsx", token: str | None = None
) -> Any:
    files = {"file": (name, data, XLSX)} if data is not None else None
    return c.post(
        PATH,
        data={"csrf_token": token if token is not None else csrf_for(c, PATH)},
        files=files,
        follow_redirects=False,
    )


def shown(response: Any) -> str:
    return text_of(response.text)


# --- the permission -----------------------------------------------------------------------------


def test_can_compare_apres_is_the_export_roles() -> None:
    for role, expected in ((ROLE_MANAGER, True), (ROLE_ACCOUNTANT, True), (ROLE_STAFF, False)):
        user = User(agency_id=1, phone="+1", role=role)
        assert permissions.can_compare_apres(user) is expected
        assert permissions.can_compare_apres(user) == permissions.can_export_day(user)
    assert not permissions.can_compare_apres(User(agency_id=1, phone="+1", role="weird"))


# --- the results --------------------------------------------------------------------------------


def test_the_form_page(closed: tuple[Env, TestClient, bytes]) -> None:
    _, acc, _ = closed
    r = acc.get(PATH)
    assert r.status_code == 200 and r.headers["cache-control"] == "no-store"
    text = shown(r)
    assert "Comparer avec la feuille de la comptable" in text
    assert f'action="{PATH}"' in r.text and 'enctype="multipart/form-data"' in r.text
    assert 'name="csrf_token"' in r.text and 'type="file"' in r.text


def test_the_same_file_has_no_difference(closed: tuple[Env, TestClient, bytes]) -> None:
    env, acc, stored = closed
    r = submit(acc, stored)
    assert r.status_code == 200 and r.headers["cache-control"] == "no-store"
    text = shown(r)
    assert "Aucune différence" in text
    assert "Seulement dans sa feuille 0" in text and "Seulement dans l'application 0" in text
    assert "Cellules différentes 0" in text and "Lignes communes 34" in text
    assert "Identiques 34" in text


def test_one_changed_amount_shows_both_values(closed: tuple[Env, TestClient, bytes]) -> None:
    _, acc, stored = closed
    r = submit(acc, edited(stored, (0, "P.B", 777)))
    text = shown(r)
    assert r.status_code == 200 and "Cellules différentes 1" in text and "Identiques 33" in text
    assert "P.B" in text and "777.00" in text
    assert fixture_tktt()[0]["DOC NUMBER"] in text


def test_one_changed_name_shows_both_names(closed: tuple[Env, TestClient, bytes]) -> None:
    _, acc, stored = closed
    r = submit(acc, edited(stored, (3, "PAX NAME", "OTHER/NAME")))
    text = shown(r)
    assert "Cellules différentes 1" in text and "OTHER/NAME" in text and "PAX NAME" in text
    assert "SYNTH" in text  # the app's value of the same cell


def test_extra_and_missing_rows(closed: tuple[Env, TestClient, bytes]) -> None:
    _, acc, stored = closed
    wb = openpyxl.load_workbook(io.BytesIO(stored))
    ws = wb.active
    assert ws is not None
    ws.delete_rows(FIRST + 5)  # the app's sixth row is missing from hers
    extra: dict[str, Any] = {
        "SEQ NO": "000999",
        "A/L": "199",
        "DOC NUMBER": "1998888888",
        "TOTAL DOC": 1234,
        "PAX NAME": "ONLY/HERS",
        "TRNC": "TKTT",
    }
    for title, value in extra.items():
        ws.cell(row=FIRST + 33, column=COLS[title]).value = value  # the free last slot
    out = io.BytesIO()
    wb.save(out)
    text = shown(submit(acc, out.getvalue()))
    assert "Seulement dans l'application 1" in text
    assert "Seulement dans sa feuille 1" in text and "ONLY/HERS" in text and "1998888888" in text


def test_the_thirteen_digit_form_still_aligns(closed: tuple[Env, TestClient, bytes]) -> None:
    _, acc, stored = closed
    line = fixture_tktt()[0]
    airline = line["A/L"]
    thirteen = airline + line["DOC NUMBER"]
    r = submit(acc, edited(stored, (0, "DOC NUMBER", thirteen)))
    text = shown(r)
    assert "Aucune différence" in text and "Lignes communes 34" in text


def test_a_changed_cell_of_text_is_normalised_before_the_comparison(
    closed: tuple[Env, TestClient, bytes],
) -> None:
    _, acc, stored = closed
    r = submit(acc, edited(stored, (0, "AG", " bruge "), (1, "AG", "BRUGE")))
    assert "Aucune différence" in shown(r)


def test_the_page_names_the_version_compared(closed: tuple[Env, TestClient, bytes]) -> None:
    _, acc, stored = closed
    assert "version 1" in shown(submit(acc, stored)).lower()


def test_a_sheet_of_another_date_is_noticed_but_compared(
    closed: tuple[Env, TestClient, bytes],
) -> None:
    _, acc, stored = closed
    wb = openpyxl.load_workbook(io.BytesIO(stored))
    ws = wb.active
    assert ws is not None
    ws["D1"] = "19OCT"
    out = io.BytesIO()
    wb.save(out)
    text = shown(submit(acc, out.getvalue()))
    assert "Attention : la feuille indique" in text and "Lignes communes 34" in text


def test_her_text_is_escaped(closed: tuple[Env, TestClient, bytes]) -> None:
    _, acc, stored = closed
    nasty = "<script>alert(1)</script>"
    r = submit(acc, edited(stored, (0, "PAX NAME", nasty), (1, "CLT", '"><img src=x onerror=1>')))
    assert nasty not in r.text and "<img src=x" not in r.text
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in r.text


def test_a_formula_in_her_file_is_never_evaluated(closed: tuple[Env, TestClient, bytes]) -> None:
    _, acc, stored = closed
    r = submit(acc, edited(stored, (0, "CLT", '=HYPERLINK("http://x","y")')))
    assert r.status_code == 200 and "http://x" not in r.text  # no cached value: an empty cell


def test_manager_may_compare_too(closed: tuple[Env, TestClient, bytes]) -> None:
    env, _, stored = closed
    assert "Aucune différence" in shown(submit(env.manager_a(), stored))


# --- what is stored -----------------------------------------------------------------------------


def test_nothing_is_stored_but_an_audit_event_with_counts(
    closed: tuple[Env, TestClient, bytes],
) -> None:
    env, acc, stored = closed
    before = (
        len(table(env, Upload)), len(table(env, AmadeusLine)), len(table(env, DayExport)),
    )  # fmt: skip
    r = submit(acc, edited(stored, (0, "P.B", 5), (2, "PAX NAME", "SECRET/PERSON")))
    assert r.status_code == 200
    assert before == (
        len(table(env, Upload)),
        len(table(env, AmadeusLine)),
        len(table(env, DayExport)),
    )
    (event,) = audit_events(env.db.owner_url, env.seeded.agency_a, audit.APRES_COMPARED)
    assert event.details["matched"] == 34 and event.details["differing_cells"] == 2
    assert event.details["only_hers"] == 0 and event.details["only_app"] == 0
    assert event.details["identical"] == 32 and event.details["version"] == 1
    flat = str(event.details)
    assert "SECRET" not in flat and "sa feuille" not in flat and "SYNTH" not in flat
    assert event.actor_user_id is not None


def test_the_file_never_goes_to_a_temporary_file(
    closed: tuple[Env, TestClient, bytes], monkeypatch: pytest.MonkeyPatch
) -> None:
    """A file above Starlette's 1 MB spool would be written out; the route keeps it in memory."""
    _, acc, stored = closed

    def no_disk(self: Any, *a: Any, **k: Any) -> None:
        raise AssertionError("an upload was written to a temporary file")

    monkeypatch.setattr(tempfile.SpooledTemporaryFile, "rollover", no_disk)
    big = add_zip_entries(stored, {"xl/media/padding.bin": os.urandom(1_400_000)})
    assert len(big) > 1_100_000
    assert "Aucune différence" in shown(submit(acc, big))


# --- the files that are refused -----------------------------------------------------------------


@pytest.mark.parametrize(
    ("data", "needle"),
    [
        (b"", "vide"),
        (b"just some text, not a workbook", ".xlsx"),
        (b"PK\x03\x04 not really a workbook", "pas pu être lu"),
    ],
)
def test_a_bad_file_is_refused_in_french(
    closed: tuple[Env, TestClient, bytes], data: bytes, needle: str
) -> None:
    _, acc, _ = closed
    r = submit(acc, data)
    assert r.status_code == 400 and needle in shown(r)


def test_no_file_at_all_is_refused(closed: tuple[Env, TestClient, bytes]) -> None:
    _, acc, _ = closed
    r = submit(acc, None)
    assert r.status_code == 400 and "Choisissez un fichier" in shown(r)


def test_an_oversized_file_is_refused_with_413(closed: tuple[Env, TestClient, bytes]) -> None:
    _, acc, _ = closed
    assert submit(acc, b"PK\x03\x04" + bytes(MAX_FILE_BYTES)).status_code == 413


def test_a_zip_bomb_like_file_is_refused(closed: tuple[Env, TestClient, bytes]) -> None:
    _, acc, stored = closed
    bomb = add_zip_entries(stored, {"xl/media/bomb.bin": bytes(25 * 1024 * 1024)})
    assert len(bomb) < 200_000
    r = submit(acc, bomb)
    assert r.status_code == 400 and "pas pu être lu" in shown(r)


def test_a_workbook_without_the_report_columns_is_refused(
    closed: tuple[Env, TestClient, bytes],
) -> None:
    _, acc, _ = closed
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = "rien"
    out = io.BytesIO()
    wb.save(out)
    r = submit(acc, out.getvalue())
    assert r.status_code == 400 and "pas pu être lu" in shown(r)


def test_a_parser_that_is_busy_gives_a_503_and_the_hit_back(
    closed: tuple[Env, TestClient, bytes], monkeypatch: pytest.MonkeyPatch
) -> None:
    env, acc, stored = closed

    def busy(_: bytes) -> Any:
        raise SandboxBusy

    env.factory.app.dependency_overrides[compare_routes.get_apres_parser] = lambda: busy  # type: ignore[attr-defined]
    monkeypatch.setattr(
        ratelimit, "APRES_COMPARE_USER", ratelimit.Limit("apres_compare_user", 1, 3600)
    )
    r = submit(acc, stored)
    assert r.status_code == 503 and r.headers["retry-after"]
    env.factory.app.dependency_overrides[compare_routes.get_apres_parser] = lambda: parse_apres  # type: ignore[attr-defined]
    assert submit(acc, stored).status_code == 200  # the busy attempt did not use the budget


def test_it_goes_through_the_real_sandbox_by_default(
    closed: tuple[Env, TestClient, bytes],
) -> None:
    env, acc, stored = closed
    del env.factory.app.dependency_overrides[compare_routes.get_apres_parser]  # type: ignore[attr-defined]
    assert "Aucune différence" in shown(submit(acc, stored))
    assert "Cellules différentes 1" in shown(submit(acc, edited(stored, (0, "P.B", 3))))


# --- who, and which days ------------------------------------------------------------------------


def test_staff_are_refused_and_nothing_is_read(
    closed: tuple[Env, TestClient, bytes],
) -> None:
    env, _, stored = closed
    spy: list[bytes] = []

    def parser(data: bytes) -> Any:
        spy.append(data)
        return parse_apres(data)

    env.factory.app.dependency_overrides[compare_routes.get_apres_parser] = lambda: parser  # type: ignore[attr-defined]
    staff = env.staff_a()
    assert staff.get(PATH).status_code == 403
    token = csrf_for(staff, "/")
    r = staff.post(PATH, data={"csrf_token": token}, files={"file": ("a.xlsx", stored, XLSX)})
    assert r.status_code == 403 and spy == []
    assert any(
        e.event_type == "permission_denied"
        for e in audit_events(env.db.owner_url, env.seeded.agency_a)
    )


def test_another_agency_gets_a_404(closed: tuple[Env, TestClient, bytes]) -> None:
    env, _, stored = closed
    other = env.manager_b()
    assert other.get(PATH).status_code == 404
    token = csrf_for(other, "/")
    r = other.post(PATH, data={"csrf_token": token}, files={"file": ("a.xlsx", stored, XLSX)})
    assert r.status_code == 404


def test_an_anonymous_visitor_is_turned_away(closed: tuple[Env, TestClient, bytes]) -> None:
    env, _, stored = closed
    anon = env.factory()
    assert anon.get(PATH, follow_redirects=False).status_code in (303, 401)
    r = anon.post(PATH, files={"file": ("a.xlsx", stored, XLSX)}, follow_redirects=False)
    assert r.status_code in (303, 401)


def test_a_day_that_is_not_closed_is_a_404(env: Env) -> None:
    upload_20oct(env)
    acc = accountant(env)
    sell_line(acc, 0)
    assert acc.get(PATH).status_code == 404
    token = csrf_for(acc, f"/days/{D20}")
    r = acc.post(PATH, data={"csrf_token": token}, files={"file": ("a.xlsx", b"x", XLSX)})
    assert r.status_code == 404


def test_a_closed_day_without_a_stored_export_is_a_404(env: Env, gate_open: None) -> None:
    m = env.manager_a()
    enter(m, tk(), staff_sign=SIGN, day="2026-10-18", pb="24000")  # no Amadeus report that day
    assert close(m, "2026-10-18").status_code == 303
    assert table(env, DayExport) == []
    path = "/days/2026-10-18/compare"
    assert m.get(path).status_code == 404
    token = csrf_for(m, "/days/2026-10-18")
    r = m.post(path, data={"csrf_token": token}, files={"file": ("a.xlsx", b"x", XLSX)})
    assert r.status_code == 404


def test_a_day_that_does_not_exist_or_is_not_a_date(closed: tuple[Env, TestClient, bytes]) -> None:
    _, acc, _ = closed
    assert acc.get("/days/2026-10-18/compare").status_code == 404
    assert acc.get("/days/not-a-date/compare").status_code == 400
    assert acc.get("/days/2099-01-01/compare").status_code == 400


def test_after_a_reopen_the_page_is_a_404_until_it_is_closed_again(
    closed: tuple[Env, TestClient, bytes],
) -> None:
    env, acc, _ = closed
    from tests.test_sale_pages import post

    m = env.manager_a()
    assert (
        post(
            m, f"/days/{D20}/reopen", {"reason": "prix corrigé"}, via=f"/days/{D20}/reopen"
        ).status_code
        == 303
    )
    assert acc.get(PATH).status_code == 404


def test_the_day_page_offers_it_only_when_it_can_work(
    closed: tuple[Env, TestClient, bytes],
) -> None:
    env, acc, _ = closed
    link = f'href="{PATH}"'
    assert link in acc.get(f"/days/{D20}").text
    assert "Comparer avec la feuille de la comptable" in shown(acc.get(f"/days/{D20}"))
    assert link not in env.staff_a().get(f"/days/{D20}").text


def test_the_day_page_does_not_offer_it_for_an_open_day(env: Env) -> None:
    upload_20oct(env)
    acc = accountant(env)
    sell_line(acc, 0)
    assert f'href="{PATH}"' not in acc.get(f"/days/{D20}").text


# --- csrf, origin, rate limit -------------------------------------------------------------------


def test_csrf_and_origin_are_checked_before_the_body_is_read(
    closed: tuple[Env, TestClient, bytes],
) -> None:
    env, acc, stored = closed
    spy: list[bytes] = []

    def parser(data: bytes) -> Any:
        spy.append(data)
        return parse_apres(data)

    env.factory.app.dependency_overrides[compare_routes.get_apres_parser] = lambda: parser  # type: ignore[attr-defined]
    files = {"file": ("a.xlsx", stored, XLSX)}
    assert acc.post(PATH, files=files).status_code == 403  # no token
    assert acc.post(PATH, data={"csrf_token": "x" * 43}, files=files).status_code == 403
    token = csrf_for(acc, PATH)
    foreign = acc.post(
        PATH, data={"csrf_token": token}, files=files, headers={"Origin": "https://evil.example"}
    )
    assert foreign.status_code == 403
    header = acc.post(PATH, files=files, headers={"X-CSRF-Token": token})
    assert header.status_code == 200 and len(spy) == 1  # a header token works like the field
    assert acc.post(PATH, files=files, headers={"X-CSRF-Token": "y" * 43}).status_code == 403
    assert len(spy) == 1


def test_there_is_a_rate_limit_and_a_wrong_token_still_counts(
    closed: tuple[Env, TestClient, bytes], monkeypatch: pytest.MonkeyPatch
) -> None:
    """A bad form token is not refunded: otherwise one user could push unmetered 10 MB bodies."""
    env, acc, stored = closed
    monkeypatch.setattr(
        ratelimit, "APRES_COMPARE_USER", ratelimit.Limit("apres_compare_user", 2, 3600)
    )
    token = csrf_for(acc, PATH)
    assert submit(acc, stored, token="z" * 43).status_code == 403  # counted
    assert submit(acc, stored, token="z" * 43).status_code == 403  # counted
    r = submit(acc, stored, token=token)
    assert r.status_code == 429 and r.headers["retry-after"]
    assert submit(env.manager_a(), stored).status_code == 200  # a budget of its own


def test_a_stored_export_that_cannot_be_read_back_is_a_french_page_and_a_refund(
    closed: tuple[Env, TestClient, bytes], monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.parsing.amadeus import ParseError

    env, acc, stored = closed
    monkeypatch.setattr(
        ratelimit, "APRES_COMPARE_USER", ratelimit.Limit("apres_compare_user", 1, 3600)
    )

    def broken(_: bytes) -> Any:
        raise ParseError("secret detail")

    monkeypatch.setattr(compare_routes, "parse_apres", broken)
    r = submit(acc, stored)
    assert r.status_code == 503 and "enregistré" in shown(r) and "secret detail" not in r.text
    monkeypatch.setattr(compare_routes, "parse_apres", parse_apres)
    assert submit(acc, stored).status_code == 200  # the server fault did not use the budget


# --- the number of statements -------------------------------------------------------------------

CEILING = 25


def test_the_number_of_statements_does_not_grow_with_the_differences(
    closed: tuple[Env, TestClient, bytes],
) -> None:
    env, acc, stored = closed
    engine = env.factory.app.state.engine  # type: ignore[attr-defined]
    token = csrf_for(acc, PATH)
    assert acc.get(PATH).status_code == 200  # warm-up
    with counting(engine) as form:
        assert acc.get(PATH).status_code == 200
    with counting(engine) as same:
        assert submit(acc, stored, token=token).status_code == 200
    many = edited(stored, *[(i, "CLT", f"autre {i}") for i in range(30)])
    with counting(engine) as different:
        r = submit(acc, many, token=token)
    assert "Cellules différentes 30" in shown(r)
    assert len(form) <= 12 and len(same) <= CEILING
    assert len(different) == len(same)

"""Reading the accountant's APRES: the four hand-typed columns, read only on request.

An Amadeus upload never carries P.B / AG / CLT / AP into the app (``parse_report`` leaves them
unread); the parallel-run comparison asks for them with ``parse_apres``, which runs in the same
sandbox as every other upload. Synthetic data only.
"""

from __future__ import annotations

import json
from decimal import Decimal
from typing import Any

import openpyxl
import pytest

from app.parsing.amadeus import MANUAL_LIMITS, ParseError, parse_apres, parse_report
from app.parsing.sandbox import parse_in_sandbox
from app.parsing.wire import decode_message, encode_report
from tests.builders import AMADEUS_HEADERS, build_workbook, ticket_row

APRES_HEADERS = [*AMADEUS_HEADERS, "P.B", "AG", "CLT", "AP"]


def apres_book(rows: list[list[Any]]) -> bytes:
    return build_workbook(rows, headers=APRES_HEADERS)


def row(
    pb: Any = 1500, ag: Any = "BRUGE", clt: Any = "Client X", ap: Any = "+22212345678"
) -> list[Any]:
    return [*ticket_row(), pb, ag, clt, ap]


def test_apres_reads_the_four_hand_typed_columns() -> None:
    (ticket,) = parse_apres(apres_book([row()])).tickets
    assert (ticket.pb, ticket.ag, ticket.clt, ticket.ap) == (
        Decimal("1500.00"),
        "BRUGE",
        "Client X",
        "+22212345678",
    )


def test_an_amadeus_upload_never_reads_them() -> None:
    (ticket,) = parse_report(apres_book([row()])).tickets
    assert (ticket.pb, ticket.ag, ticket.clt, ticket.ap) == (None, None, None, None)


def test_empty_cells_are_none_and_a_zero_is_a_zero() -> None:
    blank, zero = parse_apres(
        apres_book([row(None, None, None, None), [*ticket_row(doc="3940000002"), 0, "", " ", ""]])
    ).tickets
    assert (blank.pb, blank.ag, blank.clt, blank.ap) == (None, None, None, None)
    assert zero.pb == Decimal("0.00") and (zero.ag, zero.clt, zero.ap) == (None, None, None)


def test_an_unreadable_amount_is_ignored_with_a_warning() -> None:
    report = parse_apres(apres_book([row(pb="beaucoup")]))
    assert report.tickets[0].pb is None
    assert any("P.B" in w for w in report.warnings)


def test_a_formula_cell_is_read_as_its_cached_value_never_evaluated() -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws["C1"], ws["D1"], ws["C3"], ws["D3"] = (
        "Report date range",
        "01OCT",
        "Current date",
        "02-OCT-2026",
    )
    for col, title in enumerate(APRES_HEADERS, start=1):
        ws.cell(row=5, column=col, value=title)
    for col, value in enumerate([*ticket_row(), "=1+1", "=HYPERLINK(1)", None, None], start=1):
        ws.cell(row=6, column=col, value=value)
    import io

    buf = io.BytesIO()
    wb.save(buf)
    (ticket,) = parse_apres(buf.getvalue()).tickets
    assert ticket.pb is None and ticket.ag is None  # no cached value: nothing is computed


def test_the_hand_typed_text_is_bounded() -> None:
    for key in ("ag", "clt", "ap"):
        assert MANUAL_LIMITS[key] >= 100
    long = "X" * (MANUAL_LIMITS["clt"] + 50)
    report = parse_apres(apres_book([row(clt=long)]))
    assert report.tickets[0].clt == long[: MANUAL_LIMITS["clt"]]
    assert any("CLT" in w for w in report.warnings)


def test_the_wire_carries_them_and_checks_them() -> None:
    report = parse_apres(apres_book([row()]))
    again = decode_message(encode_report(report))
    assert (again.tickets[0].pb, again.tickets[0].ag) == (Decimal("1500.00"), "BRUGE")
    message = json.loads(encode_report(report))
    for key, bad in (("pb", "1e5"), ("pb", "12345678901234"), ("ag", "x" * 500), ("clt", "a\x00b"),
                     ("ap", 5)):  # fmt: skip
        broken = json.loads(json.dumps(message))
        broken["report"]["tickets"][0][key] = bad
        with pytest.raises(ParseError):
            decode_message(json.dumps(broken).encode())
    del message["report"]["tickets"][0]["pb"]
    with pytest.raises(ParseError):
        decode_message(json.dumps(message).encode())


def test_it_runs_in_the_sandbox_like_any_upload() -> None:
    parsed = parse_in_sandbox(
        apres_book([row()]), timeout_s=30, memory_mb=512, parse_target=parse_apres
    )
    assert parsed.tickets[0].pb == Decimal("1500.00") and parsed.tickets[0].clt == "Client X"

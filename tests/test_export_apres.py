"""The TJQ APRES export (B10): rows, workbook, formula neutralisation, storage at close, download.

Synthetic data only (tests/fixtures/synthetic_20oct.py). Service tests run on the SQLite
``session`` fixture; the pages on ``web_db``.
"""

from __future__ import annotations

import hashlib
import io
from collections.abc import Iterator
from datetime import UTC, date, datetime
from decimal import Decimal
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import approvals, export_apres, partners, permissions, sales
from app.auth import ratelimit
from app.export_apres import ApresLine, ApresReport
from app.ingest import ingest_report
from app.models import (
    AmadeusLine,
    DayExport,
    ReportDay,
    StaffSignLink,
    User,
)
from app.repository import create_agency
from app.sales import SaleInput, TicketInput
from app.tenancy import rebind_tenant
from tests.builders import amadeus_line, build_amadeus_xlsx
from tests.fixtures import synthetic_20oct as fx
from tests.page_support import Env
from tests.test_sale_pages import (
    SIGN,
    accountant,
    enter,
    post,
    table,
    tk,
)
from tests.web_support import WebDb, audit_events, owner_session

NOW = datetime(2026, 10, 21, 12, 0, tzinfo=UTC)
DAY = date(2026, 10, 20)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADERS = ["SEQ NO", "A/L", "DOC NUMBER", "TOTAL DOC", "PAX NAME", "AS", "TRNC", "RLOC",
           "P.B", "AG", "CLT", "AP"]  # fmt: skip
FIRST_ROW = 6
CR, LF = chr(13), chr(10)


# --- helpers ------------------------------------------------------------------------------------


def dec(value: object) -> Decimal:
    return Decimal(str(value))


def open_book(data: bytes) -> Any:
    wb = openpyxl.load_workbook(io.BytesIO(data))  # formulas are not evaluated: we see them
    ws = wb.active
    assert ws is not None
    return ws


def no_formulas(ws: Any) -> None:
    for row in ws.iter_rows():
        for cell in row:
            assert cell.data_type != "f", cell.coordinate
            assert not (isinstance(cell.value, str) and cell.value.startswith("=")), cell.coordinate


def data_rows(ws: Any, count: int) -> list[list[Any]]:
    return [
        [ws.cell(row=FIRST_ROW + i, column=c).value for c in range(1, 13)] for i in range(count)
    ]


class World:
    def __init__(self, session: Session) -> None:
        self.db = session
        self.agency_id = create_agency(session, "agence-x", "Agence X").id
        self.manager = self._user("+22212380001", "manager")
        self.accountant = self._user("+22212380002", "accountant")
        self.staff = self._user("+22212380003", "staff", "0004SV")
        other = create_agency(session, "agence-y", "Agence Y").id
        self.other_manager = self._user("+22212380004", "manager", agency_id=other)
        rebind_tenant(session, self.agency_id)
        self.partner = partners.create_partner(session, self.manager, "BADI", now=NOW)
        partners.set_fee(
            session, self.partner, Decimal("500"), date(2026, 1, 1), self.manager, now=NOW
        )

    def _user(self, phone: str, role: str, *signs: str, agency_id: int | None = None) -> User:
        user = User(agency_id=agency_id or self.agency_id, phone=phone, role=role)
        self.db.add(user)
        self.db.commit()
        for sign in signs:
            self.db.add(StaffSignLink(agency_id=user.agency_id, user_id=user.id, staff_sign=sign))
        self.db.commit()
        return user

    def upload(self, data: bytes | None = None) -> None:
        ingest_report(
            self.db, self.agency_id, data or fx.build_synthetic_20oct(), "20OCT.xlsx",
            uploaded_by=self.manager.id,
        )  # fmt: skip

    def lines(self) -> list[AmadeusLine]:
        return list(
            self.db.scalars(
                select(AmadeusLine)
                .where(AmadeusLine.agency_id == self.agency_id, AmadeusLine.trnc == "TKTT")
                .order_by(AmadeusLine.seq_order)
            )
        )

    def ticket_of(self, line: AmadeusLine) -> TicketInput:
        return TicketInput(
            doc_number=line.doc_number,
            airline=line.airline,
            total_doc=line.total_doc if line.total_doc is not None else 0,
            pax_name=line.pax_name,
            pax_type=line.pax_type,
            rloc=line.rloc,
            trnc=line.trnc,
        )

    def sell(self, *lines: AmadeusLine, partner: bool = False, **kw: Any) -> int:
        cost = sum((line.total_doc or Decimal(0)) for line in lines)
        fields: dict[str, Any] = {"staff_sign": lines[0].staff_sign, "payment_method": "cash"}
        if partner:
            fields.update(ag_kind="partner", partner_id=self.partner.id, payment_method="credit")
        else:
            fields.update(pb=cost + 1500, clt=f"Client {lines[0].seq_no}")
        fields.update(kw)
        data = SaleInput(tickets=[self.ticket_of(line) for line in lines], **fields)
        return sales.create_sale(self.db, self.manager, data, day=DAY, now=NOW).sale.id

    def sell_all(self, *, skip: frozenset[int] = frozenset()) -> dict[str, int]:
        """A sale per TKTT line, except ``skip`` (indexes): a layover pair at 10-11, partner sales
        every 7th line from the 4th."""
        lines = self.lines()
        assert len(lines) == fx.TKTT_COUNT
        made: dict[str, int] = {}
        i = 0
        while i < len(lines):
            if i in skip:
                i += 1
                continue
            if i == 10:
                made["layover"] = self.sell(lines[10], lines[11], ap="+22212345010")
                i += 2
                continue
            made[lines[i].doc_number] = self.sell(lines[i], partner=(i % 7 == 3))
            i += 1
        return made


@pytest.fixture(autouse=True)
def gate_open(monkeypatch: pytest.MonkeyPatch) -> None:
    """These tests are about storing the export; the close gate (B9) is switched off here and
    tested for real in tests/test_close_gate.py and tests/test_approval_pages.py."""
    monkeypatch.setattr(approvals, "can_close", lambda *args, **kwargs: [])


@pytest.fixture
def w(session: Session) -> World:
    return World(session)


def fixture_tktt() -> list[dict[str, Any]]:
    return [line for line in fx.LINES if line["TRNC"] == "TKTT"]


# --- acceptance on the synthetic 20 Oct ---------------------------------------------------------


def test_acceptance_workbook_of_the_synthetic_20_october(w: World) -> None:
    w.upload()
    made = w.sell_all()
    report = export_apres.apres_rows(w.db, w.manager, DAY)
    assert report is not None
    ws = open_book(export_apres.build_apres_xlsx(report))
    # rows 1-3 and the header row
    assert [ws["A1"].value, ws["C1"].value, ws["D1"].value] == [
        "Agy no",
        "Report date range",
        "20OCT",
    ]
    assert [ws["C3"].value, ws["D3"].value] == ["Current date", "21-OCT-2026"]
    assert [ws.cell(row=5, column=c).value for c in range(1, 13)] == HEADERS
    rows = data_rows(ws, 34)
    assert ws.cell(row=FIRST_ROW + 34, column=1).value is None  # nothing but the totals below
    # CANX removed, SEQ order, the 8 Amadeus columns equal to the fixture
    expected = fixture_tktt()
    assert len(expected) == 34
    for row, line in zip(rows, expected, strict=True):
        assert row[0] == line["SEQ NO"]
        assert row[1] == line["A/L"]
        assert row[2] == line["DOC NUMBER"]
        assert dec(row[3]) == Decimal(line["TOTAL DOC"])
        assert row[4] == line["PAX NAME"]
        assert row[5] == line["AS"]
        assert row[6] == "TKTT"
        assert row[7] == line["RLOC"]
    seqs = [r[0] for r in rows]
    assert seqs == sorted(seqs)
    assert sum(dec(r[3]) for r in rows) == Decimal(fx.TKTT_TOTAL) == Decimal(624_126)
    # P.B: full price on the first ticket of the layover sale, 0 on its second
    lines = w.lines()
    first, second = rows[10], rows[11]
    assert first[2] == lines[10].doc_number and second[2] == lines[11].doc_number
    layover_cost = (lines[10].total_doc or 0) + (lines[11].total_doc or 0)
    assert dec(first[8]) == Decimal(layover_cost) + 1500
    assert dec(second[8]) == 0 and isinstance(second[8], int | float)
    assert first[11] == second[11] == "+22212345010"  # a strict E.164 phone stays plain
    # AG / CLT / AP filled; partner lines carry the partner name
    for n, row in enumerate(rows):
        if n in (10, 11):  # the layover sale is a direct one
            continue
        assert row[9] == ("BADI" if n % 7 == 3 else "BRUGE"), n
        if n % 7 == 3:
            assert dec(row[8]) == Decimal(expected[n]["TOTAL DOC"]) + 500
        else:
            assert row[10] == f"Client {expected[n]['SEQ NO']}"
    # the totals row, Decimal-exact when read back
    total_row = FIRST_ROW + 34 + 1
    assert ws.cell(row=total_row, column=1).value == "TOTAUX"
    assert ws.cell(row=total_row, column=3).value == 34
    assert dec(ws.cell(row=total_row, column=4).value) == Decimal(624_126)
    pb_sum = sum(dec(r[8]) for r in rows)
    assert dec(ws.cell(row=total_row, column=9).value) == pb_sum
    assert dec(ws.cell(row=total_row, column=11).value) == pb_sum - Decimal(624_126)
    assert (report.ticket_count, report.total_doc, report.total_pb) == (34, Decimal(624126), pb_sum)
    assert report.total_margin == pb_sum - Decimal(624_126)
    assert not report.unmatched
    # amounts are numbers, 2 decimals, and nothing is a formula
    for r in range(FIRST_ROW, FIRST_ROW + 34):
        for c in (4, 9):
            cell = ws.cell(row=r, column=c)
            assert isinstance(cell.value, int | float) and cell.number_format == "#,##0.00"
    no_formulas(ws)
    # the group marker of the 2-ticket sale
    marker = f"Vente #{made['layover']} · 2 billets"
    assert [g.marker for g in report.lines[10:12]] == [marker, marker]
    assert ws.cell(row=FIRST_ROW + 10, column=1).comment.text == marker
    assert report.lines[0].marker == ""


def test_the_header_values_come_from_the_upload(w: World) -> None:
    w.upload()
    report = export_apres.apres_rows(w.db, w.manager, DAY)
    assert report is not None
    assert (report.agy_no, report.range_raw, report.current_date) == ("", "20OCT", "21-OCT-2026")
    assert report.day == DAY


def test_no_report_for_the_day_means_no_export(w: World) -> None:
    assert export_apres.apres_rows(w.db, w.manager, DAY) is None


def test_the_line_set_survives_a_second_identical_upload(w: World) -> None:
    w.upload()
    w.upload()  # identical bytes: the second upload holds no lines of its own
    report = export_apres.apres_rows(w.db, w.manager, DAY)
    assert report is not None and report.ticket_count == 34


def test_another_agency_sees_nothing_and_staff_may_not_export(w: World) -> None:
    w.upload()
    rebind_tenant(w.db, w.other_manager.agency_id)
    assert export_apres.apres_rows(w.db, w.other_manager, DAY) is None
    rebind_tenant(w.db, w.agency_id)
    with pytest.raises(PermissionError):
        export_apres.apres_rows(w.db, w.staff, DAY)
    assert export_apres.apres_rows(w.db, w.accountant, DAY) is not None


def test_a_line_without_an_app_sale_is_marked_and_empty(w: World) -> None:
    w.upload()
    w.sell_all(skip=frozenset({5}))
    report = export_apres.apres_rows(w.db, w.manager, DAY)
    assert report is not None
    gap = report.lines[5]
    assert gap.sale_id is None and gap.pb is None
    assert (gap.ag, gap.clt, gap.ap) == ("", "", "")
    assert report.unmatched == 1 and report.ticket_count == 34
    ws = open_book(export_apres.build_apres_xlsx(report))
    row = FIRST_ROW + 5
    assert [ws.cell(row=row, column=c).value for c in (9, 10, 11, 12)] == [None] * 4
    assert ws.cell(row=row, column=9).comment is not None
    assert "Aucune vente" in ws.cell(row=row, column=9).comment.text
    neighbour = ws.cell(row=row + 1, column=1)
    assert ws.cell(row=row, column=1).fill.fgColor.rgb != neighbour.fill.fgColor.rgb
    # P.B totals only count what the app holds
    assert report.total_pb == sum((ln.pb or Decimal(0)) for ln in report.lines)


def test_a_voided_sale_leaves_its_tickets_unmatched(w: World) -> None:
    w.upload()
    made = w.sell_all()
    sales.void_sale(w.db, w.manager, made["layover"], "saisie en double", now=NOW)
    report = export_apres.apres_rows(w.db, w.manager, DAY)
    assert report is not None and report.unmatched == 2
    assert report.lines[10].sale_id is None and report.lines[11].sale_id is None


def test_canx_is_dropped_even_when_written_in_lowercase(w: World) -> None:
    w.upload()
    lines = [
        ApresLine(  # what the filter keeps and drops is decided on the normalised TRNC
            seq_no="1",
            airline="199",
            doc_number="1990000001",
            total_doc=Decimal(1),
            pax_name="A/B",
            staff_sign="0004SV",
            trnc=trnc,
            rloc="AAAAAA",
        )  # fmt: skip
        for trnc in ("TKTT", " canx ", "Canx")
    ]
    kept = export_apres.without_canx(lines)
    assert [ln.trnc for ln in kept] == ["TKTT"]


# --- formula injection --------------------------------------------------------------------------


def hostile_report(text: str) -> ApresReport:
    line = ApresLine(
        seq_no=text, airline=text, doc_number="1990000001", total_doc=Decimal("10.50"),
        pax_name=text, staff_sign=text, trnc="TKTT", rloc=text,
        pb=Decimal("12"), ag=text, clt=text, ap=text, sale_id=1, ticket_count=1,
    )  # fmt: skip
    return export_apres.make_report(
        day=DAY, agy_no=text, range_raw=text, current_date=text, lines=[line]
    )


@pytest.mark.parametrize("lead", ["=", "+", "-", "@", "\t", "\r"])
def test_text_starting_with_a_formula_character_is_stored_as_text(lead: str) -> None:
    text = f'{lead}HYPERLINK("http://x","y")'
    ws = open_book(export_apres.build_apres_xlsx(hostile_report(text)))
    expected = "'" + text.replace(CR, LF)  # XML reads a CR back as a line feed
    for col in (1, 2, 5, 6, 8, 10, 11, 12):  # every text column of the data row
        cell = ws.cell(row=FIRST_ROW, column=col)
        assert cell.value.replace(CR, LF) == expected, col
        assert cell.data_type == "s", col
    for coordinate in ("B1", "D1", "D3"):  # the header block
        assert ws[coordinate].value.replace(CR, LF) == expected
        assert ws[coordinate].data_type == "s"
    no_formulas(ws)
    assert ws.cell(row=FIRST_ROW, column=4).value == 10.5  # the amounts stay numbers


def test_ordinary_text_is_not_touched_and_illegal_characters_do_not_crash() -> None:
    assert export_apres.neutralise("SYNTH01/PAX") == "SYNTH01/PAX"
    assert export_apres.neutralise("") == ""
    assert export_apres.neutralise("a=b") == "a=b"
    assert export_apres.neutralise("x\x00y\x07") == "xy"
    ws = open_book(export_apres.build_apres_xlsx(hostile_report("bad\x00name")))
    assert ws.cell(row=FIRST_ROW, column=5).value == "badname"


def test_a_hostile_client_name_from_a_real_sale_is_neutralised(w: World) -> None:
    w.upload()
    line = w.lines()[0]
    w.sell(line, clt='=cmd|" /C calc"!A0')
    report = export_apres.apres_rows(w.db, w.manager, DAY)
    assert report is not None
    ws = open_book(export_apres.build_apres_xlsx(report))
    cell = ws.cell(row=FIRST_ROW, column=11)
    assert cell.value == '\'=cmd|" /C calc"!A0' and cell.data_type == "s"
    no_formulas(ws)


# --- the corrections hook (filled by B9) --------------------------------------------------------


def test_corrections_render_highlighted_cells_with_comments() -> None:
    report = hostile_report("x")
    report = export_apres.make_report(
        day=DAY, agy_no="", range_raw="20OCT", current_date="21-OCT-2026", lines=list(report.lines)
    )
    doc = report.lines[0].doc_number
    corrections = {
        (doc, "TOTAL DOC"): ("1000", "1200.50", "Comptable A"),
        (doc, "PAX NAME"): ("OLD/NAME", "NEW/NAME", "Manager A"),
    }
    plain = open_book(export_apres.build_apres_xlsx(report))
    ws = open_book(export_apres.build_apres_xlsx(report, corrections=corrections))
    amount, name = ws.cell(row=FIRST_ROW, column=4), ws.cell(row=FIRST_ROW, column=5)
    assert dec(amount.value) == Decimal("1200.50") and name.value == "NEW/NAME"
    assert "1000" in amount.comment.text and "Comptable A" in amount.comment.text
    assert "OLD/NAME" in name.comment.text and "Manager A" in name.comment.text
    untouched = ws.cell(row=FIRST_ROW, column=8)
    assert untouched.comment is None
    assert amount.fill.fgColor.rgb != plain.cell(row=FIRST_ROW, column=4).fill.fgColor.rgb
    no_formulas(ws)


def test_by_default_there_are_no_corrections() -> None:
    assert export_apres.NO_CORRECTIONS == {}
    report = hostile_report("x")
    assert export_apres.build_apres_xlsx(report) == export_apres.build_apres_xlsx(
        report, corrections={}
    )


# --- the stored export at close -----------------------------------------------------------------


def test_closing_through_the_service_stores_version_1_then_2(w: World) -> None:
    w.upload()
    w.sell_all()
    report_day = w.db.scalars(select(ReportDay)).one()
    first = export_apres.close_day_with_export(w.db, w.accountant, report_day, now=NOW)
    assert first is not None and first.version == 1
    assert first.sha256 == hashlib.sha256(first.content).hexdigest()
    assert first.created_by_user_id == w.accountant.id
    assert len(open_book(first.content)["A"]) > 34
    from app import days

    days.reopen_day(w.db, w.manager, report_day, "correction du prix", now=NOW)
    second = export_apres.close_day_with_export(w.db, w.accountant, report_day, now=NOW)
    assert second is not None and second.version == 2
    assert export_apres.stored_versions(w.db, w.agency_id, report_day.id) == [1, 2]
    assert export_apres.get_export(w.db, w.agency_id, report_day.id, None).version == 2  # type: ignore[union-attr]
    assert export_apres.get_export(w.db, w.agency_id, report_day.id, 1).version == 1  # type: ignore[union-attr]
    assert export_apres.get_export(w.db, w.agency_id, report_day.id, 3) is None


def test_closing_without_a_report_stores_nothing(w: World) -> None:
    from app import days

    report_day = days.get_or_open_day(w.db, w.manager, DAY, now=NOW)
    stored = export_apres.close_day_with_export(w.db, w.manager, report_day, now=NOW)
    assert stored is None and report_day.status == "closed"
    assert w.db.scalars(select(DayExport)).all() == []


def test_the_export_of_another_agency_cannot_be_read(w: World) -> None:
    w.upload()
    w.sell_all()
    report_day = w.db.scalars(select(ReportDay)).one()
    export_apres.close_day_with_export(w.db, w.manager, report_day, now=NOW)
    assert export_apres.get_export(w.db, w.other_manager.agency_id, report_day.id, None) is None
    assert export_apres.stored_versions(w.db, w.other_manager.agency_id, report_day.id) == []


def test_the_filename_is_built_from_the_date_only() -> None:
    assert export_apres.export_filename(date(2026, 10, 20)) == "TJQ_APRES_2026-10-20.xlsx"
    assert (
        export_apres.export_filename(date(2026, 10, 20), draft=True)
        == "TJQ_APRES_2026-10-20_BROUILLON.xlsx"
    )


def test_can_export_day_is_manager_and_accountant_only(w: World) -> None:
    assert permissions.can_export_day(w.manager) and permissions.can_export_day(w.accountant)
    assert not permissions.can_export_day(w.staff)
    assert not permissions.can_export_day(User(agency_id=1, phone="+1", role="weird"))


# --- the pages ----------------------------------------------------------------------------------


@pytest.fixture
def env(web_db: WebDb) -> Env:
    e = Env(web_db)
    e.link(e.seeded.staff_a, SIGN)
    return e


D20 = "2026-10-20"


def upload_20oct(env: Env) -> None:
    env.clock.now = datetime(2026, 10, 21, 12, 0, tzinfo=UTC)
    with owner_session(env.db.owner_url) as s:
        ingest_report(
            s, env.seeded.agency_a, fx.build_synthetic_20oct(), "20OCT.xlsx",
            uploaded_by=env.seeded.manager_a,
        )  # fmt: skip


def sell_line(c: TestClient, index: int, **fields: Any) -> int:
    line = fixture_tktt()[index]
    ticket = tk(
        line["DOC NUMBER"], str(line["TOTAL DOC"]), airline=line["A/L"], pax_name=line["PAX NAME"],
        rloc=line["RLOC"],
    )  # fmt: skip
    return enter(
        c, ticket, staff_sign=line["AS"], day=D20, pb=str(line["TOTAL DOC"] + 1000), **fields
    )


def close(c: TestClient, day: str = D20) -> Any:
    return post(c, f"/days/{day}/close", {}, via=f"/days/{day}/close")


def exports(env: Env) -> list[DayExport]:
    return sorted(table(env, DayExport), key=lambda e: e.version)


@pytest.fixture
def closed_day(env: Env) -> Iterator[tuple[Env, TestClient]]:
    upload_20oct(env)
    acc = accountant(env)
    sell_line(acc, 0)
    sell_line(acc, 1)
    assert close(acc).status_code == 303
    yield env, acc


def test_closing_through_the_page_stores_version_1(closed_day: tuple[Env, TestClient]) -> None:
    env, acc = closed_day
    (stored,) = exports(env)
    assert stored.version == 1 and stored.agency_id == env.seeded.agency_a
    assert stored.sha256 == hashlib.sha256(stored.content).hexdigest()
    ws = open_book(stored.content)
    assert ws.cell(row=FIRST_ROW, column=3).value == fixture_tktt()[0]["DOC NUMBER"]
    assert ws.cell(row=FIRST_ROW, column=10).value == "BRUGE"
    assert ws.cell(row=FIRST_ROW + 2, column=10).value is None  # no sale for the third line
    kinds = [e.event_type for e in audit_events(env.db.owner_url, env.seeded.agency_a)]
    assert "day_export_stored" in kinds


def test_reopen_and_close_again_stores_version_2_and_v1_stays_downloadable(
    closed_day: tuple[Env, TestClient],
) -> None:
    env, acc = closed_day
    m = env.manager_a()
    r = post(m, f"/days/{D20}/reopen", {"reason": "prix corrigé"}, via=f"/days/{D20}/reopen")
    assert r.status_code == 303
    sell_line(m, 2)  # a sale added while reopened
    assert close(acc).status_code == 303
    v1, v2 = exports(env)
    assert (v1.version, v2.version) == (1, 2) and v1.content != v2.content
    latest = acc.get(f"/days/{D20}/apres.xlsx")
    assert latest.status_code == 200 and latest.content == v2.content
    old = acc.get(f"/days/{D20}/apres.xlsx?v=1")
    assert old.status_code == 200 and old.content == v1.content
    assert acc.get(f"/days/{D20}/apres.xlsx?v=2").content == v2.content
    assert acc.get(f"/days/{D20}/apres.xlsx?v=3").status_code == 404
    page = acc.get(f"/days/{D20}").text
    assert f"/days/{D20}/apres.xlsx?v=1" in page and f"/days/{D20}/apres.xlsx?v=2" in page


def test_download_headers_and_filename(closed_day: tuple[Env, TestClient]) -> None:
    env, acc = closed_day
    for c in (acc, env.manager_a()):
        r = c.get(f"/days/{D20}/apres.xlsx")
        assert r.status_code == 200
        assert r.headers["content-type"] == XLSX
        assert (
            r.headers["content-disposition"] == 'attachment; filename="TJQ_APRES_2026-10-20.xlsx"'
        )
        assert r.headers["cache-control"] == "no-store"
        assert r.headers["x-content-type-options"] == "nosniff"
        assert r.content[:2] == b"PK"


@pytest.mark.parametrize("v", ["abc", "0", "-1", "1.5", "1 ", "٣", "99999999999999999999", ""])
def test_an_invalid_version_is_a_400(closed_day: tuple[Env, TestClient], v: str) -> None:
    _, acc = closed_day
    r = acc.get(f"/days/{D20}/apres.xlsx", params={"v": v})
    assert r.status_code == 400
    assert "Version invalide" in r.text and r.headers["cache-control"] == "no-store"


def test_an_open_day_has_no_stored_export(env: Env) -> None:
    upload_20oct(env)
    m = env.manager_a()
    sell_line(m, 0)
    r = m.get(f"/days/{D20}/apres.xlsx")
    assert r.status_code == 404 and "Journée non clôturée" in r.text
    assert m.get(f"/days/{D20}/apres.xlsx?v=1").status_code == 404
    assert m.get("/days/2026-10-19/apres.xlsx").status_code == 404  # no such day at all


def test_a_bad_date_is_a_400(env: Env) -> None:
    m = env.manager_a()
    for raw in ("2026-13-45", "x", "2030-01-01"):
        assert m.get(f"/days/{raw}/apres.xlsx").status_code == 400
        assert m.get(f"/days/{raw}/apres-preview.xlsx").status_code == 400


def test_staff_are_refused_and_audited(closed_day: tuple[Env, TestClient]) -> None:
    env, _ = closed_day
    s = env.staff_a()
    for path in (f"/days/{D20}/apres.xlsx", f"/days/{D20}/apres.xlsx?v=1",
                 f"/days/{D20}/apres-preview.xlsx"):  # fmt: skip
        assert s.get(path).status_code == 403
    denied = audit_events(env.db.owner_url, env.seeded.agency_a, "permission_denied")
    assert len(denied) >= 3
    assert not audit_events(env.db.owner_url, env.seeded.agency_a, "export_downloaded")
    assert "apres" not in s.get(f"/days/{D20}").text


def test_another_agency_gets_a_404(closed_day: tuple[Env, TestClient]) -> None:
    env, _ = closed_day
    b = env.manager_b()
    assert b.get(f"/days/{D20}/apres.xlsx").status_code == 404
    assert b.get(f"/days/{D20}/apres.xlsx?v=1").status_code == 404
    assert b.get(f"/days/{D20}/apres-preview.xlsx").status_code == 404  # no report there


def test_downloads_are_audited_with_ids_only(closed_day: tuple[Env, TestClient]) -> None:
    env, acc = closed_day
    assert acc.get(f"/days/{D20}/apres.xlsx?v=1").status_code == 200
    (event,) = audit_events(env.db.owner_url, env.seeded.agency_a, "export_downloaded")
    (day_row,) = table(env, ReportDay)
    assert event.details == {"report_day_id": day_row.id, "version": 1, "draft": False}
    assert event.actor_user_id is not None


def test_the_preview_of_an_open_day_is_a_draft_and_is_not_stored(env: Env) -> None:
    upload_20oct(env)
    m = env.manager_a()
    sell_line(m, 0)
    r = m.get(f"/days/{D20}/apres-preview.xlsx")
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    assert r.headers["content-disposition"] == (
        'attachment; filename="TJQ_APRES_2026-10-20_BROUILLON.xlsx"'
    )
    assert (
        r.headers["cache-control"] == "no-store"
        and r.headers["x-content-type-options"] == "nosniff"
    )
    ws = open_book(r.content)
    assert ws.cell(row=FIRST_ROW, column=10).value == "BRUGE"
    assert exports(env) == []
    (event,) = audit_events(env.db.owner_url, env.seeded.agency_a, "export_downloaded")
    assert event.details["draft"] is True
    acc = accountant(env)
    assert acc.get(f"/days/{D20}/apres-preview.xlsx").status_code == 200


def test_the_preview_needs_a_report_and_an_open_day(closed_day: tuple[Env, TestClient]) -> None:
    env, acc = closed_day
    assert (
        acc.get(f"/days/{D20}/apres-preview.xlsx").status_code == 409
    )  # closed: use the stored one
    assert acc.get("/days/2026-10-17/apres-preview.xlsx").status_code == 404  # no report


def test_closing_a_day_without_a_report_stores_nothing_and_says_so(env: Env) -> None:
    m = env.manager_a()
    enter(m, tk(), staff_sign=SIGN, day="2026-10-18", pb="24000")
    assert close(m, "2026-10-18").status_code == 303
    assert exports(env) == []
    page = m.get("/days/2026-10-18").text
    assert "Aucun export enregistré" in page
    assert m.get("/days/2026-10-18/apres.xlsx").status_code == 404


def test_the_day_page_buttons(env: Env) -> None:
    upload_20oct(env)
    m, acc = env.manager_a(), accountant(env)
    sell_line(m, 0)
    for c in (m, acc):
        page = c.get(f"/days/{D20}").text
        assert f'href="/days/{D20}/apres-preview.xlsx"' in page
        assert f"/days/{D20}/apres.xlsx" not in page  # nothing stored yet
    assert close(acc).status_code == 303
    for c in (m, acc):
        page = c.get(f"/days/{D20}").text
        assert f'href="/days/{D20}/apres.xlsx"' in page
        assert "apres-preview" not in page
    staff = env.staff_a().get(f"/days/{D20}").text
    assert "apres" not in staff


def test_closing_twice_does_not_store_a_second_version(closed_day: tuple[Env, TestClient]) -> None:
    env, acc = closed_day
    assert post(acc, f"/days/{D20}/close", {}, via=f"/days/{D20}").status_code == 409
    assert [e.version for e in exports(env)] == [1]


def test_a_failed_export_rolls_the_close_back(env: Env, monkeypatch: pytest.MonkeyPatch) -> None:
    upload_20oct(env)
    acc = accountant(env)
    sell_line(acc, 0)

    def boom(*args: Any, **kwargs: Any) -> bytes:
        raise RuntimeError("disk full")

    monkeypatch.setattr(export_apres, "build_apres_xlsx", boom)
    assert close(acc).status_code == 500  # the error page; nothing was saved
    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, env.seeded.agency_a)
        assert s.scalars(select(ReportDay.status)).all() == ["open"]
        assert s.scalars(select(DayExport)).all() == []


# --- P.B never vanishes, duplicates, bounds, neutralising (review round) -----------------------

DOC_A, DOC_B = "1990000001", "1990000002"


def custom_report(*specs: tuple[str, str]) -> bytes:
    """A made-up report of ``(doc, trnc)`` lines, in order."""
    lines = [
        amadeus_line(seq=f"{n:06d}", doc=doc, trnc=trnc, rloc=f"AB{n:04d}", total=1000)
        for n, (doc, trnc) in enumerate(specs, start=1)
    ]
    return build_amadeus_xlsx(lines, range_raw="20OCT", current="21-OCT-2026")


def layover(w: World, pb: str = "5000") -> int:
    tickets = [
        TicketInput(d, "199", "1000", "SYNTH01/PAX", "ADT", "ABC123", "TKTT")
        for d in (DOC_A, DOC_B)
    ]
    data = SaleInput(
        tickets=tickets, pb=pb, clt="Client", payment_method="cash", staff_sign="0004SV"
    )
    return sales.create_sale(w.db, w.manager, data, day=DAY, now=NOW).sale.id


def test_pb_moves_to_the_first_exported_line_when_the_first_ticket_is_canx(w: World) -> None:
    w.upload(custom_report((DOC_A, "CANX"), (DOC_B, "TKTT")))
    layover(w)
    report = export_apres.apres_rows(w.db, w.manager, DAY)
    assert report is not None and [ln.doc_number for ln in report.lines] == [DOC_B]
    assert report.lines[0].pb == Decimal(5000) and report.total_pb == Decimal(5000)


def test_pb_moves_to_the_first_line_present_when_the_first_ticket_is_missing(w: World) -> None:
    w.upload(custom_report(("1990000009", "TKTT"), (DOC_B, "TKTT")))
    layover(w)
    report = export_apres.apres_rows(w.db, w.manager, DAY)
    assert report is not None
    assert [ln.pb for ln in report.lines] == [None, Decimal(5000)]


def test_pb_follows_the_report_order_not_the_entry_order(w: World) -> None:
    w.upload(custom_report((DOC_B, "TKTT"), (DOC_A, "TKTT")))
    layover(w)
    report = export_apres.apres_rows(w.db, w.manager, DAY)
    assert report is not None
    assert [ln.pb for ln in report.lines] == [Decimal(5000), Decimal(0)]


def test_a_sale_with_no_exported_ticket_puts_no_pb(w: World) -> None:
    w.upload(custom_report((DOC_A, "CANX"), (DOC_B, "CANX"), ("1990000003", "TKTT")))
    layover(w)
    report = export_apres.apres_rows(w.db, w.manager, DAY)
    assert report is not None and report.total_pb == 0 and report.unmatched == 1


def test_a_second_line_with_the_same_number_is_a_duplicate(w: World) -> None:
    w.upload()
    w.sell_all()
    first = w.lines()[0]
    last = max(ln.seq_order for ln in w.db.scalars(select(AmadeusLine)))
    w.db.add(
        AmadeusLine(
            agency_id=w.agency_id,
            upload_id=first.upload_id,
            seq_no="999999",
            seq_order=last + 1,
            doc_number="199" + first.doc_number,  # the same ticket with its airline prefix
            airline=first.airline,
            total_doc=first.total_doc,
            pax_name=first.pax_name,
            pax_type=first.pax_type,
            staff_sign=first.staff_sign,
            rloc=first.rloc,
            trnc="TKTT",
            flags=[],
            source_row=last + 1,
        )
    )
    w.db.commit()
    report = export_apres.apres_rows(w.db, w.manager, DAY)
    assert report is not None and report.ticket_count == 35
    dup = report.lines[-1]
    assert dup.duplicate and dup.pb is None and (dup.ag, dup.clt, dup.ap) == ("", "", "")
    assert report.lines[0].pb is not None and not report.lines[0].duplicate
    ws = open_book(export_apres.build_apres_xlsx(report))
    cell = ws.cell(row=FIRST_ROW + 34, column=9)
    assert "Doublon" in cell.comment.text and cell.value is None
    assert ws.cell(row=FIRST_ROW + 34, column=1).fill.fgColor.rgb != "00000000"


def test_closing_is_refused_after_50_versions_and_changes_nothing(w: World) -> None:
    from app import days

    w.upload()
    report_day = days.get_or_open_day(w.db, w.manager, DAY, now=NOW)
    for n in range(1, export_apres.MAX_VERSIONS + 1):
        w.db.add(
            DayExport(
                agency_id=w.agency_id, report_day_id=report_day.id, version=n, content=b"x",
                sha256="a" * 64, created_by_user_id=w.manager.id,
            )
        )  # fmt: skip
    w.db.commit()
    with pytest.raises(export_apres.ExportError, match="Trop de versions"):
        export_apres.close_day_with_export(w.db, w.accountant, report_day, now=NOW)
    assert report_day.status == "open"
    assert export_apres.stored_versions(w.db, w.agency_id, report_day.id)[-1] == 50


def test_a_report_beyond_the_line_cap_or_a_workbook_beyond_the_size_cap_is_refused(
    w: World, monkeypatch: pytest.MonkeyPatch
) -> None:
    w.upload()
    monkeypatch.setattr(export_apres, "MAX_LINES", 10)
    with pytest.raises(export_apres.ExportError, match="trop long"):
        export_apres.apres_rows(w.db, w.manager, DAY)
    monkeypatch.setattr(export_apres, "MAX_LINES", 50_000)
    monkeypatch.setattr(export_apres, "MAX_BYTES", 10)
    from app import days

    report_day = days.get_or_open_day(w.db, w.manager, DAY, now=NOW)
    with pytest.raises(export_apres.ExportError, match="volumineux"):
        export_apres.store_next_version(w.db, w.manager, report_day, now=NOW)


def test_the_preview_of_a_too_long_report_is_a_409(
    env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    upload_20oct(env)
    monkeypatch.setattr(export_apres, "MAX_LINES", 10)
    r = env.manager_a().get(f"/days/{D20}/apres-preview.xlsx")
    assert r.status_code == 409 and "trop long" in r.text


@pytest.mark.parametrize("phone", ["+22212345678", "+12345678", "+123456789012345"])
def test_a_strict_e164_phone_stays_plain_text(phone: str) -> None:
    assert export_apres.neutralise(phone) == phone


@pytest.mark.parametrize(
    "text",
    ["+1234567", "+1234567890123456", "+222 12345678", "+22212345678x", "+٢٢٢١٢٣٤٥٦٧٨",
     " =1+1", "\n=1+1", "\x0b-1", "\x0c@a", "\xa0=1", "\xa0 +x", "\n"],
)  # fmt: skip
def test_everything_else_that_looks_like_a_formula_is_quoted(text: str) -> None:
    expected = "'" + text.replace(chr(11), "").replace(chr(12), "")  # VT and FF are not XML
    assert export_apres.neutralise(text) == expected


def test_blank_or_ordinary_text_is_left_alone() -> None:
    for text in ("  ", "a +b", "x=1", "12+3"):
        assert export_apres.neutralise(text) == text


def test_a_corrected_cell_is_always_written_and_its_comment_is_clean() -> None:
    report = hostile_report("x")
    doc = report.lines[0].doc_number
    corrections = {(doc, "PAX NAME"): ("OLD\x00/NAME", "", "Man\x07ager")}
    ws = open_book(export_apres.build_apres_xlsx(report, corrections=corrections))
    cell = ws.cell(row=FIRST_ROW, column=5)
    assert cell.value in (None, "") and cell.comment is not None
    assert "OLD/NAME" in cell.comment.text and "Manager" in cell.comment.text
    assert cell.fill.fgColor.rgb != "00000000"


def test_downloads_have_a_light_per_user_limit(
    closed_day: tuple[Env, TestClient], monkeypatch: pytest.MonkeyPatch
) -> None:
    env, acc = closed_day
    monkeypatch.setattr(ratelimit, "EXPORT_USER", ratelimit.Limit("export_user", 2, 3600))
    assert acc.get(f"/days/{D20}/apres.xlsx").status_code == 200
    assert acc.get(f"/days/{D20}/apres.xlsx").status_code == 200
    assert acc.get(f"/days/{D20}/apres.xlsx").status_code == 429
    assert acc.get(f"/days/{D20}/apres-preview.xlsx").status_code == 429
    assert env.manager_a().get(f"/days/{D20}/apres.xlsx").status_code == 200  # own budget


def test_a_compared_day_exports_the_compared_upload_not_the_newest(w: World) -> None:
    from app import days
    from app.models import Upload

    w.upload()
    first = w.db.scalars(select(Upload).order_by(Upload.id)).first()
    assert first is not None
    w.upload(custom_report((DOC_A, "TKTT")))  # a newer upload for the same date
    newest = export_apres.apres_rows(w.db, w.manager, DAY)
    assert newest is not None and newest.ticket_count == 1
    report_day = days.get_or_open_day(w.db, w.manager, DAY, now=NOW)
    report_day.compared_upload_id = first.id
    w.db.commit()
    compared = export_apres.apres_rows(w.db, w.manager, DAY)
    assert compared is not None and compared.ticket_count == 34

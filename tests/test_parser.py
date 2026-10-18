from __future__ import annotations

import io
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from app.parsing import amadeus
from app.parsing.amadeus import ParsedReport, ParseError, parse_report
from tests.builders import (
    AMADEUS_HEADERS,
    build_workbook,
    set_cached_formula,
    ticket_row,
)

FIXTURE = Path(__file__).parent / "fixtures" / "1OCT.xlsx"


@pytest.fixture(scope="module")
def report() -> ParsedReport:
    return parse_report(FIXTURE)


def by_name(report: ParsedReport, name: str) -> list[amadeus.ParsedTicket]:
    return [t for t in report.tickets if t.passenger_name == name]


# --- acceptance on the real fixture -------------------------------------------------


def test_fixture_counts_and_total(report: ParsedReport) -> None:
    assert len(report.tickets) == 48
    assert sum(t.cost for t in report.tickets if t.cost is not None) == Decimal(1118857)
    assert report.warnings == ()
    assert all(t.flags == () for t in report.tickets)


def test_fixture_header(report: ParsedReport) -> None:
    from datetime import date

    assert report.header.report_range_raw == "01OCT"
    assert report.header.report_date == date(2026, 10, 1)
    assert report.header.generated_on == date(2026, 10, 2)


def test_fixture_pax_types(report: ParsedReport) -> None:
    inf = {t.passenger_name for t in report.tickets if t.pax_type == "INF"}
    chd = {t.passenger_name for t in report.tickets if t.pax_type == "CHD"}
    assert inf == {"PAX02/GIVEN02", "PAX37/GIVEN37"}
    assert chd == {"PAX03/GIVEN03"}
    assert sum(t.pax_type == "ADT" for t in report.tickets) == 45


def test_fixture_first_ticket_keeps_text_and_zeros(report: ParsedReport) -> None:
    first = report.tickets[0]
    assert first.seq_no == "081327"
    assert first.airline == "199"
    assert first.ticket_no == "9990000110"
    assert first.cost == Decimal(24050)
    assert isinstance(first.cost, Decimal)
    assert first.passenger_name == "PAX23/GIVEN23 MRS"
    assert first.staff_sign == "0004SV"
    assert first.transaction_type == "TKTT"
    assert first.booking_ref == "QA0009"
    assert first.source_row == 6


def test_fixture_repeat_passenger(report: ParsedReport) -> None:
    assert len(by_name(report, "PAX35/GIVEN35")) == 2


def test_fixture_ticket_numbers_unique(report: ParsedReport) -> None:
    assert len({t.ticket_no for t in report.tickets}) == 48


# --- columns I..N must have zero effect ---------------------------------------------


def _variant(tmp_path: Path, name: str, mutate) -> Path:  # type: ignore[no-untyped-def]
    wb = openpyxl.load_workbook(FIXTURE, data_only=False)
    ws = wb.active
    assert ws is not None
    mutate(ws)
    path = tmp_path / name
    wb.save(path)
    return path


def _clear(ws) -> None:  # type: ignore[no-untyped-def]
    for row in ws.iter_rows(min_col=9, max_col=14):
        for cell in row:
            cell.value = None


def _junk(ws) -> None:  # type: ignore[no-untyped-def]
    formulas = ["=SUM(D6:D53)", "=1/0", "=NOPE()", "garbage", 12.5, "=D6*2"]
    for r in range(1, 60):
        for i, col in enumerate(range(9, 15)):
            ws.cell(row=r, column=col, value=formulas[(r + i) % len(formulas)])


def _delete(ws) -> None:  # type: ignore[no-untyped-def]
    ws.delete_cols(9, 6)


@pytest.mark.parametrize("mutate", [_clear, _junk, _delete], ids=["cleared", "junk", "deleted"])
def test_hand_typed_columns_have_no_effect(
    tmp_path: Path,
    report: ParsedReport,
    mutate,  # type: ignore[no-untyped-def]
) -> None:
    other = parse_report(_variant(tmp_path, "v.xlsx", mutate))
    assert other.header == report.header
    assert other.tickets == report.tickets
    assert other.warnings == report.warnings


# --- input types and limits ---------------------------------------------------------


def test_accepts_bytes_and_file_object(report: ParsedReport) -> None:
    data = FIXTURE.read_bytes()
    assert parse_report(data) == report
    assert parse_report(io.BytesIO(data)) == report
    with FIXTURE.open("rb") as fh:
        assert parse_report(fh) == report


def test_rejects_oversize_bytes() -> None:
    with pytest.raises(ParseError, match="10 MB"):
        parse_report(b"0" * (10 * 1024 * 1024 + 1))


def test_rejects_oversize_stream() -> None:
    with pytest.raises(ParseError, match="10 MB"):
        parse_report(io.BytesIO(b"0" * (10 * 1024 * 1024 + 1)))


def test_rejects_oversize_path(tmp_path: Path) -> None:
    big = tmp_path / "big.xlsx"
    big.write_bytes(b"0" * (10 * 1024 * 1024 + 1))
    with pytest.raises(ParseError, match="10 MB"):
        parse_report(big)


def test_rejects_missing_path(tmp_path: Path) -> None:
    with pytest.raises(ParseError, match="not found"):
        parse_report(tmp_path / "nope.xlsx")


@pytest.mark.parametrize("data", [b"", b"not a zip at all", b"PK\x03\x04garbage"])
def test_rejects_corrupt_or_non_xlsx(data: bytes) -> None:
    with pytest.raises(ParseError):
        parse_report(data)


def test_rejects_csv_renamed_to_xlsx(tmp_path: Path) -> None:
    path = tmp_path / "x.xlsx"
    path.write_text("SEQ NO,A/L\n1,2\n")
    with pytest.raises(ParseError):
        parse_report(path)


# --- headers ------------------------------------------------------------------------


def test_missing_required_header_names_it() -> None:
    headers = [h for h in AMADEUS_HEADERS if h != "RLOC"]
    data = build_workbook([], headers=headers)
    with pytest.raises(ParseError, match="RLOC"):
        parse_report(data)


def test_no_doc_number_header_means_no_header_row() -> None:
    data = build_workbook([], headers=["A", "B", "C"])
    with pytest.raises(ParseError, match="DOC NUMBER"):
        parse_report(data)


def test_header_row_found_wherever_it_is() -> None:
    data = build_workbook([ticket_row()], header_row=9)
    result = parse_report(data)
    assert len(result.tickets) == 1
    assert result.tickets[0].source_row == 10


def test_columns_mapped_by_name_not_position() -> None:
    headers = ["RLOC", "TRNC", "AS", "PAX NAME", "TOTAL DOC", "DOC NUMBER", "A/L", "SEQ NO"]
    row = list(reversed(ticket_row(seq="000009", rloc="QWERTY", cost=777)))
    result = parse_report(build_workbook([row], headers=headers))
    t = result.tickets[0]
    assert (t.seq_no, t.booking_ref, t.cost) == ("000009", "QWERTY", Decimal(777))


def test_header_names_are_case_and_space_tolerant() -> None:
    headers = [" seq no", "a/l", "Doc Number ", "total doc", "pax name", "as", "trnc", "rloc"]
    assert len(parse_report(build_workbook([ticket_row()], headers=headers)).tickets) == 1


def test_missing_report_date_range_label() -> None:
    with pytest.raises(ParseError, match="Report date range"):
        parse_report(build_workbook([ticket_row()], range_raw=None))


def test_missing_current_date() -> None:
    with pytest.raises(ParseError, match="Current date"):
        parse_report(build_workbook([ticket_row()], current=None))


@pytest.mark.parametrize("bad", ["OCT1", "32OCT", "01XXX", "", "01OCT-05OCT"])
def test_unparseable_range(bad: str) -> None:
    with pytest.raises(ParseError, match="Report date range"):
        parse_report(build_workbook([ticket_row()], range_raw=bad))


@pytest.mark.parametrize("bad", ["2026-10-02", "02-OCT", "xx"])
def test_unparseable_current_date(bad: str) -> None:
    with pytest.raises(ParseError, match="Current date"):
        parse_report(build_workbook([ticket_row()], current=bad))


def test_year_rollover() -> None:
    from datetime import date

    result = parse_report(build_workbook([ticket_row()], range_raw="31DEC", current="01-JAN-2027"))
    assert result.header.report_date == date(2026, 12, 31)
    assert result.header.generated_on == date(2027, 1, 1)
    assert result.header.report_range_raw == "31DEC"


def test_same_month_keeps_year() -> None:
    from datetime import date

    result = parse_report(build_workbook([ticket_row()], range_raw="30SEP", current="30-SEP-2026"))
    assert result.header.report_date == date(2026, 9, 30)


def test_impossible_date_is_parse_error() -> None:
    # 29FEB generated in 2027 -> year 2027 is not a leap year
    with pytest.raises(ParseError, match="Report date range"):
        parse_report(build_workbook([ticket_row()], range_raw="29FEB", current="01-MAR-2027"))


# --- rows ---------------------------------------------------------------------------


def test_blank_rows_in_middle_and_end_are_skipped() -> None:
    rows = [
        ticket_row(doc="3940000001"),
        [None] * 8,
        ticket_row(doc="3940000002"),
        ["  ", None, "", None, None, None, None, None],
        ticket_row(doc="3940000003"),
        [None] * 8,
        [None] * 8,
    ]
    result = parse_report(build_workbook(rows))
    assert [t.ticket_no for t in result.tickets] == ["3940000001", "3940000002", "3940000003"]
    assert [t.source_row for t in result.tickets] == [6, 8, 10]
    assert result.warnings == ()


def test_empty_report_has_no_tickets() -> None:
    result = parse_report(build_workbook([]))
    assert result.tickets == ()


@pytest.mark.parametrize("trnc", ["RFND", "CANX", "EMDA", "EMDS", "tktt ", ""])
def test_unknown_trnc_is_flagged_not_dropped(trnc: str) -> None:
    result = parse_report(build_workbook([ticket_row(trnc=trnc or None)]))
    assert len(result.tickets) == 1
    flagged = "unknown_trnc" in result.tickets[0].flags
    assert flagged is (trnc.strip().upper() != "TKTT")


def test_trnc_is_normalised_and_known_has_no_flag() -> None:
    t = parse_report(build_workbook([ticket_row(trnc=" tktt ")])).tickets[0]
    assert t.transaction_type == "TKTT"
    assert t.flags == ()


def test_numeric_ticket_and_seq_cells_converted_with_warning() -> None:
    row = ticket_row(seq=81327, doc=9990000110.0, airline=199)
    result = parse_report(build_workbook([row]))
    t = result.tickets[0]
    assert t.ticket_no == "9990000110"
    assert t.seq_no == "81327"
    assert t.airline == "199"
    assert len(result.warnings) == 3
    assert all("row 6" in w for w in result.warnings)


def test_negative_cost_allowed() -> None:
    t = parse_report(build_workbook([ticket_row(cost=-5000)])).tickets[0]
    assert t.cost == Decimal(-5000)
    assert t.flags == ()


def test_decimal_cost_exact() -> None:
    t = parse_report(build_workbook([ticket_row(cost=1234.5)])).tickets[0]
    assert t.cost == Decimal("1234.5")


@pytest.mark.parametrize("cost", [None, "abc", "", "NaN", "Infinity", True])
def test_missing_or_bad_cost(cost: object) -> None:
    t = parse_report(build_workbook([ticket_row(cost=cost)])).tickets[0]
    assert t.cost is None
    assert "missing_cost" in t.flags


def test_text_numeric_cost_is_accepted() -> None:
    t = parse_report(build_workbook([ticket_row(cost="12,345")])).tickets[0]
    assert t.cost == Decimal(12345)


def test_formula_cost_uses_cached_value() -> None:
    data = build_workbook([ticket_row(cost="=1+1")])
    data = set_cached_formula(data, "D6", "1+1", 12345)
    t = parse_report(data).tickets[0]
    assert t.cost == Decimal(12345)
    assert t.flags == ()


def test_formula_cost_without_cache_is_missing() -> None:
    data = build_workbook([ticket_row(cost="=1+1")])
    t = parse_report(data).tickets[0]
    assert t.cost is None
    assert t.flags == ("missing_cost",)


def test_duplicate_doc_number_keeps_first_with_warning() -> None:
    rows = [
        ticket_row(doc="3940000001", cost=100),
        ticket_row(doc="3940000001", cost=999),
        ticket_row(doc="3940000002"),
    ]
    result = parse_report(build_workbook(rows))
    assert [t.cost for t in result.tickets] == [Decimal(100), Decimal(1000)]
    assert len(result.warnings) == 1
    assert "3940000001" in result.warnings[0]
    assert "row 7" in result.warnings[0]


def test_blank_doc_number_with_data_is_skipped_with_warning() -> None:
    result = parse_report(build_workbook([ticket_row(doc=None), ticket_row(doc="3940000002")]))
    assert [t.ticket_no for t in result.tickets] == ["3940000002"]
    assert len(result.warnings) == 1
    assert "row 6" in result.warnings[0]


@pytest.mark.parametrize(
    ("raw", "name", "pax"),
    [
        ("PAX02/GIVEN02        (INF)", "PAX02/GIVEN02", "INF"),
        ("PAX03/GIVEN03           (CHD)", "PAX03/GIVEN03", "CHD"),
        ("DOE/JOHN MR", "DOE/JOHN MR", "ADT"),
        ("  DOE/JOHN  ", "DOE/JOHN", "ADT"),
        ("DOE/JOHN (inf) ", "DOE/JOHN", "INF"),
        ("DOE/(CHD)JOHN", "DOE/(CHD)JOHN", "ADT"),
        ("O'BRIEN/ÉLODIE 😀", "O'BRIEN/ÉLODIE 😀", "ADT"),
    ],
)
def test_pax_type_and_name(raw: str, name: str, pax: str) -> None:
    t = parse_report(build_workbook([ticket_row(name=raw)])).tickets[0]
    assert (t.passenger_name, t.pax_type) == (name, pax)


def test_blank_name_is_kept_empty_adult() -> None:
    t = parse_report(build_workbook([ticket_row(name=None)])).tickets[0]
    assert (t.passenger_name, t.pax_type) == ("", "ADT")


def test_hand_typed_columns_in_synthetic_file_ignored() -> None:
    plain = parse_report(build_workbook([ticket_row()]))
    noisy = parse_report(build_workbook([ticket_row()], hand_typed=True))
    assert plain.tickets == noisy.tickets


def test_large_report_parses() -> None:
    rows = [ticket_row(doc=f"{3940000000 + i}", seq=f"{i:06d}") for i in range(10_000)]
    assert len(parse_report(build_workbook(rows)).tickets) == 10_000

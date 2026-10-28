"""The partner statement as Excel (phase F, slice F3): a closed period from its frozen snapshot
only, an open one « PROVISOIRE » from the live account; exact cells and totals; no formula can
ever come from typed text; manager and accountant only, agency-bound, rate-limited, audited."""

from __future__ import annotations

import io
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select

from app import agency_details, partner_book, partners, sales
from app.auth import ratelimit
from app.models import ROLE_ACCOUNTANT, PartnerPeriod, ReportDay, User
from app.sales import SaleInput, TicketInput
from app.tenancy import rebind_tenant
from tests.page_support import Env
from tests.web_support import (
    PASSWORD_S,
    T0,
    WebDb,
    add_user,
    audit_events,
    owner_session,
)

ACCOUNTANT_PHONE = "+22212345693"
OCT_1, OCT_3, OCT_16 = date(2026, 10, 1), date(2026, 10, 3), date(2026, 10, 16)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DETAILS = {
    "name_fr": "Agence Test",
    "name_ar": "وكالة تجريبية",
    "address": "Rue Test, Nouakchott",
    "phones": "+222 45 00 00 00 / 22 00 00 00",
    "email": "contact@agence-test.example",
    "nif": "00000000",
    "bank_name": "Banque Test",
    "bank_account": "00000 00000 00000000000 00",
    "bankily": "22000000",
    "masrvi": "22000001",
    "sedad": "22000002",
}
HYPERLINK = '=HYPERLINK("http://example.invalid","clic")'


@pytest.fixture
def env(web_db: WebDb) -> Env:
    return Env(web_db)


def _user(s: Any, user_id: int) -> User:
    found = s.scalars(select(User).where(User.id == user_id)).one()
    assert isinstance(found, User)
    return found


def ticket(doc: str, cost: str) -> TicketInput:
    return TicketInput(
        doc_number=doc, airline="199", total_doc=cost, pax_name="SYNTH01/PAX", pax_type="ADT",
        rloc="ABC123", trnc="TKTT",
    )  # fmt: skip


def partner_sale(partner_id: int, *tickets: TicketInput) -> SaleInput:
    return SaleInput(
        ag_kind="partner", partner_id=partner_id, payment_method="credit", staff_sign="0004SV",
        tickets=tickets,
    )  # fmt: skip


def set_days(env: Env, status: str) -> None:
    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, env.seeded.agency_a)
        for row in s.scalars(select(ReportDay).where(ReportDay.agency_id == env.seeded.agency_a)):
            row.status = status
            row.closed_at, row.closed_by_user_id = (
                (T0, env.seeded.manager_a) if status == "closed" else (None, None)
            )
        s.commit()


def save_details(env: Env, values: dict[str, str] | None = None) -> None:
    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, env.seeded.agency_a)
        agency_details.save_details(
            s,
            _user(s, env.seeded.manager_a),
            DETAILS if values is None else values,
            request_key=partner_book.new_request_key(),
            now=T0,
        )
        s.commit()


def setup_account(
    env: Env,
    *,
    name: str = "Partenaire Test",
    reason: str = "pénalité de la compagnie",
    reference: str = "VIR 77",
    agency_id: int | None = None,
) -> tuple[int, int]:
    """A partner (fee 300) with, from 1 to 15 October: a sale of 21 368 on the 3rd, an ADM of
    1 500 on the 10th and a cash payment of 10 000 on the 12th. ``(partner id, sale id)``."""
    agency = agency_id or env.seeded.agency_a
    manager_id = env.seeded.manager_a if agency_id is None else env.seeded.manager_b
    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, agency)
        manager = _user(s, manager_id)
        partner = partners.create_partner(s, manager, name, now=T0)
        partners.set_fee(s, partner, Decimal("300"), date(2026, 1, 1), manager, now=T0)
        s.commit()
        doc = "3940000001" if agency_id is None else "3940000002"
        result = sales.create_sale(
            s, manager, partner_sale(partner.id, ticket(doc, "21368")), day=OCT_3, now=T0
        )
        s.commit()
        partner_book.add_line(
            s, manager, partner.id, kind="ADM", amount="1500", reason=reason,
            request_key=partner_book.new_request_key(), entry_date="2026-10-10", now=T0,
        )  # fmt: skip
        partner_book.record_payment(
            s, manager, partner.id, amount="10000", method="cash", reference=reference,
            received_by_user_id=manager.id, request_key=partner_book.new_request_key(),
            paid_on="2026-10-12", now=T0,
        )  # fmt: skip
        s.commit()
        return partner.id, result.sale.id


def close_first_period(env: Env, partner_id: int) -> None:
    set_days(env, "closed")
    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, env.seeded.agency_a)
        partner_book.close_period(
            s, _user(s, env.seeded.manager_a), partner_id, OCT_1,
            request_key=partner_book.new_request_key(), now=T0,
        )  # fmt: skip
        s.commit()


def edit_sale_after_close(env: Env, partner_id: int, sale_id: int, cost: str) -> None:
    set_days(env, "open")
    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, env.seeded.agency_a)
        sales.edit_sale(
            s, _user(s, env.seeded.manager_a), sale_id,
            partner_sale(partner_id, ticket("3940000001", cost)), "coût corrigé après envoi",
            now=T0,
        )  # fmt: skip
        s.commit()


def url(partner_id: int, start: date = OCT_1) -> str:
    return f"/partner-accounts/{partner_id}/periods/{start.isoformat()}/statement.xlsx"


def workbook(content: bytes) -> Worksheet:
    wb = openpyxl.load_workbook(io.BytesIO(content))
    ws = wb.active
    assert ws is not None and len(wb.worksheets) == 1
    return ws


def _plain(value: object) -> object:
    if isinstance(value, bool):
        return value
    if isinstance(value, int | float | Decimal):
        return Decimal(str(value))
    return value


def rows(ws: Worksheet) -> list[tuple[object, ...]]:
    """Every row of the sheet, trailing empty cells dropped (an empty row is ``()``)."""
    out: list[tuple[object, ...]] = []
    for row in ws.iter_rows(values_only=True):
        values = [_plain(v) for v in row]
        while values and values[-1] is None:
            values.pop()
        out.append(tuple(values))
    while out and out[-1] == ():
        out.pop()
    return out


def D(text: str) -> Decimal:  # noqa: N802 - short on purpose, used in the expected tables
    return Decimal(text)


HEADER = [
    ("Agence Test",),
    ("وكالة تجريبية",),
    ("Rue Test, Nouakchott",),
    ("Tél. : +222 45 00 00 00 / 22 00 00 00 · E-mail : contact@agence-test.example",),
    ("NIF : 00000000",),
    ("Banque : Banque Test · RIB : 00000 00000 00000000000 00",),
    ("Bankily : 22000000 · Masrvi : 22000001 · Sedad : 22000002",),
    (),
]
TICKET_HEAD = ("Date", "N° billet", "Compagnie", "Passager", "PNR", "Coût", "Frais", "P.T")
LINE_HEAD = ("Date", "Type", "Motif", None, None, None, "État", "Montant")
PAYMENT_HEAD = ("Date", "Mode", "Référence", None, None, None, None, "Montant")
SIGN = (
    "Solde positif : le partenaire doit à l'agence ; solde négatif : l'agence doit au partenaire.",
)


def closed_statement_rows(name: str = "Partenaire Test") -> list[tuple[object, ...]]:
    return [
        *HEADER,
        ("RELEVÉ DE COMPTE PARTENAIRE",),
        ("Partenaire", name),
        ("Période", "du 01/10/2026 au 15/10/2026"),
        ("État", "Clôturée le 18/10/2026"),
        ("Édité le", "18/10/2026"),
        ("Devise", "MRU"),
        (),
        TICKET_HEAD,
        (datetime(2026, 10, 3), "3940000001", "Tunisair", "SYNTH01/PAX", "ABC123",
         D("21368"), D("300"), D("21668")),
        ("Total des billets", None, None, None, None, D("21368"), D("300"), D("21668")),
        (),
        ("Lignes manuelles",),
        LINE_HEAD,
        (datetime(2026, 10, 10), "ADM (pénalité de la compagnie)", "pénalité de la compagnie",
         None, None, None, "comptée", D("1500")),
        ("Total des lignes comptées", None, None, None, None, None, None, D("1500")),
        (),
        ("Paiements",),
        PAYMENT_HEAD,
        (datetime(2026, 10, 12), "Espèces", "VIR 77", None, None, None, None, D("10000")),
        ("Total des paiements", None, None, None, None, None, None, D("10000")),
        (),
        ("Récapitulatif",),
        ("Solde précédent", None, None, None, None, None, None, D("0")),
        ("Billets (+)", None, None, None, None, None, None, D("21668")),
        ("Lignes manuelles (+/−)", None, None, None, None, None, None, D("1500")),
        ("Paiements (−)", None, None, None, None, None, None, D("10000")),
        ("Nouveau solde", None, None, None, None, None, None, D("13168")),
        ("Sens du solde", "Le partenaire doit ce montant à l'agence."),
        SIGN,
    ]  # fmt: skip


# --- the closed period: exact cells, from the snapshot -----------------------------------------


def test_a_closed_period_gives_exact_cells_and_totals(env: Env) -> None:
    save_details(env)
    pid, _ = setup_account(env)
    close_first_period(env, pid)
    response = env.manager_a().get(url(pid))
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    assert response.headers["cache-control"] == "no-store"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["content-disposition"] == (
        f'attachment; filename="releve-partenaire-{pid}-2026-10-01-2026-10-15.xlsx"'
    )
    ws = workbook(response.content)
    assert rows(ws) == closed_statement_rows()
    assert ws.title == "Relevé"
    # Nothing is frozen: a frozen pane under the ticket header would pin the whole agency
    # header block (16 rows) on screen too (review, 27 Oct).
    assert ws.freeze_panes is None
    assert ws["A16"].value == "Date" and ws["A16"].font.bold
    # A closed statement never says « PROVISOIRE », anywhere.
    assert not any("PROVISOIRE" in str(v) for r in rows(ws) for v in r)
    # Amounts are numbers with two decimals, dates real dates shown dd/mm/yyyy.
    assert ws["F17"].data_type == "n" and ws["F17"].number_format == "#,##0.00"
    assert ws["H35"].value == Decimal("13168.00") and ws["H35"].number_format == "#,##0.00"
    assert ws["A35"].value == "Nouveau solde" and ws["H35"].font.bold
    assert ws["A17"].is_date and ws["A17"].number_format == "dd/mm/yyyy"
    for letter in "ABCDEFGH":
        assert (ws.column_dimensions[letter].width or 0) >= 10, letter


def test_the_closed_statement_never_changes_when_a_sale_changes_after_the_close(env: Env) -> None:
    save_details(env)
    pid, sale_id = setup_account(env)
    close_first_period(env, pid)
    m = env.manager_a()
    before = rows(workbook(m.get(url(pid)).content))
    edit_sale_after_close(env, pid, sale_id, "30000")
    after = rows(workbook(m.get(url(pid)).content))
    assert after == before == closed_statement_rows()
    # The difference went to the open period, as a correction (its statement is provisional).
    open_rows = rows(workbook(m.get(url(pid, OCT_16)).content))
    assert ("Solde précédent", None, None, None, None, None, None, D("13168")) in open_rows
    assert ("Nouveau solde", None, None, None, None, None, None, D("21800")) in open_rows
    correction = next(r for r in open_rows if len(r) > 1 and r[1] == "Correction après clôture")
    assert correction[0] == datetime(2026, 10, 18) and correction[-1] == D("8632")
    assert str(correction[2]).startswith(f"Vente #{sale_id} modifiée après la clôture")


def test_the_closed_statement_reads_the_snapshot_not_the_live_sales(
    env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The live ticket lines of a closed period are never read: a sale line source that would
    say otherwise is not even asked for that period."""
    save_details(env)
    pid, _ = setup_account(env)
    close_first_period(env, pid)
    from app import partner_statement

    def refuse(*args: Any, **kwargs: Any) -> Any:
        raise AssertionError("a closed statement must not read the live account")

    monkeypatch.setattr(partner_statement, "_from_live", refuse)
    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, env.seeded.agency_a)
        statement = partner_statement.statement_for(
            s, _user(s, env.seeded.manager_a), pid, OCT_1, today=T0.date()
        )
    assert not statement.provisional and statement.closing == Decimal("13168.00")
    assert rows(workbook(partner_statement.build_statement_xlsx(statement))) == (
        closed_statement_rows()
    )


def test_a_snapshot_that_does_not_add_up_is_refused(env: Env) -> None:
    from app import partner_statement

    save_details(env)
    pid, _ = setup_account(env)
    close_first_period(env, pid)
    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, env.seeded.agency_a)
        row = s.scalars(select(PartnerPeriod).where(PartnerPeriod.partner_id == pid)).one()
        snapshot = dict(row.snapshot or {})
        snapshot["sales"] = []  # the lines no longer give the frozen ticket total
        with pytest.raises(partner_statement.StatementError):
            partner_statement.from_snapshot(row, snapshot)


# --- the open period: « PROVISOIRE » --------------------------------------------------------------


def test_an_open_period_is_marked_provisoire_in_the_header_and_the_file_name(env: Env) -> None:
    pid, _ = setup_account(env)
    response = env.manager_a().get(url(pid))
    assert response.status_code == 200
    assert response.headers["content-disposition"] == (
        f'attachment; filename="releve-partenaire-{pid}-2026-10-01-2026-10-15-PROVISOIRE.xlsx"'
    )
    found = rows(workbook(response.content))
    assert ("RELEVÉ DE COMPTE PARTENAIRE — PROVISOIRE",) in found
    assert (
        "État",
        "PROVISOIRE : période ouverte, les montants peuvent encore changer",
    ) in found
    assert ("Nouveau solde", None, None, None, None, None, None, D("13168")) in found


def test_an_open_period_shows_a_pending_adjustment_as_not_counted(env: Env) -> None:
    pid, _ = setup_account(env)
    acc_id = add_user(
        env.db.owner_url, env.seeded.agency_a, ACCOUNTANT_PHONE, role=ROLE_ACCOUNTANT,
        display_name="Comptable A",
    )  # fmt: skip
    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, env.seeded.agency_a)
        partner_book.add_line(
            s, _user(s, acc_id), pid, kind="AJUSTEMENT", amount="200", direction="moins",
            reason="geste commercial", request_key=partner_book.new_request_key(),
            entry_date="2026-10-11", now=T0,
        )  # fmt: skip
        s.commit()
    found = rows(workbook(env.manager_a().get(url(pid)).content))
    assert (
        datetime(2026, 10, 11), "Ajustement", "geste commercial", None, None, None,
        "en attente (non comptée)", D("-200"),
    ) in found  # fmt: skip
    assert ("Nouveau solde", None, None, None, None, None, None, D("13168")) in found


def test_without_details_the_header_is_the_agency_name_and_nothing_else(env: Env) -> None:
    pid, _ = setup_account(env)
    found = rows(workbook(env.manager_a().get(url(pid)).content))
    assert found[0] == ("Agence A",)  # the agency's own name (``agencies.name``)
    assert found[1] == ()
    assert not any(str(r[0]).startswith(("NIF", "Tél.", "Bankily")) for r in found if r)


# --- formula injection --------------------------------------------------------------------------


def _formula_cells(content: bytes) -> list[str]:
    ws = workbook(content)
    return [c.coordinate for row in ws.iter_rows() for c in row if c.data_type == "f"]


def test_typed_text_can_never_become_a_formula(env: Env) -> None:
    save_details(env, {**DETAILS, "name_fr": "=1+1", "address": "@SUM(A1:A2)", "nif": "-2+3"})
    pid, _ = setup_account(env, name=HYPERLINK, reason="+1+1 frais", reference="=cmd|' /C calc'!A0")
    close_first_period(env, pid)
    for start in (OCT_1, OCT_16):
        content = env.manager_a().get(url(pid, start)).content
        assert _formula_cells(content) == []
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            sheet = archive.read("xl/worksheets/sheet1.xml").decode()
        assert "<f>" not in sheet and "<f " not in sheet
    ws = workbook(env.manager_a().get(url(pid)).content)
    texts = {c.value: c for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
    for typed in (HYPERLINK, "=1+1", "@SUM(A1:A2)", "+1+1 frais", "=cmd|' /C calc'!A0"):
        cell = texts.get("'" + typed)
        assert cell is not None, typed
        assert cell.data_type == "s", typed
        assert typed not in texts  # never written bare
    assert "NIF : -2+3" in texts  # a labelled header line cannot start a formula


# --- who may, isolation, limits and audit -------------------------------------------------------


def test_the_route_needs_the_statement_permission() -> None:
    from fastapi.routing import APIRoute

    from app.web.routes import partner_statements as routes

    [route] = [r for r in routes.router.routes if isinstance(r, APIRoute)]
    calls = [d.call for d in route.dependant.dependencies]
    assert routes.require_statement_reader in calls
    assert routes.require_statement_reader.__name__ == "require_can_export_partner_statement"


def test_the_accountant_downloads_it(env: Env) -> None:
    pid, _ = setup_account(env)
    add_user(
        env.db.owner_url, env.seeded.agency_a, ACCOUNTANT_PHONE, role=ROLE_ACCOUNTANT,
        display_name="Comptable A",
    )  # fmt: skip
    response = env.as_user(ACCOUNTANT_PHONE, PASSWORD_S).get(url(pid))
    assert response.status_code == 200 and response.headers["content-type"] == XLSX


def test_staff_get_403_and_an_audit_event(env: Env) -> None:
    pid, _ = setup_account(env)
    response = env.staff_a().get(url(pid))
    assert response.status_code == 403
    assert response.headers["content-type"].startswith("text/html")
    (event,) = audit_events(env.db.owner_url, env.seeded.agency_a, "permission_denied")
    assert event.actor_user_id == env.seeded.staff_a and event.details["path"] == url(pid)
    assert audit_events(env.db.owner_url, env.seeded.agency_a, "partner_statement_downloaded") == []


def test_another_agency_gets_a_404(env: Env) -> None:
    pid, _ = setup_account(env)
    response = env.manager_b().get(url(pid))
    assert response.status_code == 404
    assert "Partenaire Test" not in response.text
    assert audit_events(env.db.owner_url, env.seeded.agency_b, "partner_statement_downloaded") == []


@pytest.mark.parametrize("start", ["2026-10-02", "2026-09-16", "2030-01-01", "2026-13-01"])
def test_a_period_that_does_not_exist_is_a_404_or_400(env: Env, start: str) -> None:
    pid, _ = setup_account(env)
    response = env.manager_a().get(f"/partner-accounts/{pid}/periods/{start}/statement.xlsx")
    assert response.status_code in (400, 404)


def test_a_download_is_audited_with_ids_only(env: Env) -> None:
    pid, _ = setup_account(env, name="Partenaire Secret")
    close_first_period(env, pid)
    m = env.manager_a()
    assert m.get(url(pid)).status_code == 200
    assert m.get(url(pid, OCT_16)).status_code == 200
    events = audit_events(env.db.owner_url, env.seeded.agency_a, "partner_statement_downloaded")
    period_id = events[0].details["period_id"]
    assert isinstance(period_id, int)
    assert [e.details for e in events] == [
        {"partner_id": pid, "period_start": "2026-10-01", "period_id": period_id,
         "provisional": False},
        {"partner_id": pid, "period_start": "2026-10-16", "period_id": None, "provisional": True},
    ]  # fmt: skip
    assert all(e.actor_user_id == env.seeded.manager_a for e in events)
    assert "Secret" not in repr([e.details for e in events])


def test_downloads_are_rate_limited(env: Env, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        ratelimit, "PARTNER_STATEMENT_USER", ratelimit.Limit("partner_statement_user", 2, 3600)
    )
    pid, _ = setup_account(env)
    m = env.manager_a()
    assert [m.get(url(pid)).status_code for _ in range(3)] == [200, 200, 429]
    (event,) = audit_events(env.db.owner_url, env.seeded.agency_a, "rate_limited")
    assert event.details == {"scope": "partner_statement_user"}


def test_the_file_name_never_carries_the_partner_name(env: Env) -> None:
    pid, _ = setup_account(env, name="Partenaire Test")
    disposition = env.manager_a().get(url(pid)).headers["content-disposition"]
    assert re.fullmatch(
        r'attachment; filename="releve-partenaire-\d+-\d{4}-\d{2}-\d{2}-\d{4}-\d{2}-\d{2}'
        r'(-PROVISOIRE)?\.xlsx"',
        disposition,
    )
    assert "Partenaire" not in disposition


# --- the buttons ----------------------------------------------------------------------------------


def _page_ok(html: str) -> None:
    assert not re.search(r"\sstyle\s*=", html, re.I) and "<style" not in html.lower()
    assert not re.search(r"<script(?![^>]*\ssrc=)", html, re.I)
    assert not re.search(r"\son[a-z]+\s*=", html, re.I)


def test_the_period_and_account_pages_offer_the_download(env: Env) -> None:
    pid, _ = setup_account(env)
    close_first_period(env, pid)
    m = env.manager_a()
    period = m.get(f"/partner-accounts/{pid}/periods/2026-10-01").text
    assert "Télécharger le relevé (Excel)" in period and f'href="{url(pid)}"' in period
    _page_ok(period)
    account = m.get(f"/partner-accounts/{pid}").text
    assert "Télécharger le relevé (Excel)" in account
    assert f'href="{url(pid, OCT_16)}"' in account  # the current (open) period
    _page_ok(account)


def test_the_accountant_sees_the_buttons_too(env: Env) -> None:
    pid, _ = setup_account(env)
    add_user(
        env.db.owner_url, env.seeded.agency_a, ACCOUNTANT_PHONE, role=ROLE_ACCOUNTANT,
        display_name="Comptable A",
    )  # fmt: skip
    c: TestClient = env.as_user(ACCOUNTANT_PHONE, PASSWORD_S)
    assert f'href="{url(pid)}"' in c.get(f"/partner-accounts/{pid}").text
    assert f'href="{url(pid)}"' in c.get(f"/partner-accounts/{pid}/periods/2026-10-01").text


# --- review fixes (27 Oct) ------------------------------------------------------------------------


def dt(day: int, month: int = 10) -> datetime:
    return datetime(2026, month, day)


def total(label: str, amount: str) -> tuple[object, ...]:
    return (label, None, None, None, None, None, None, D(amount))


def new_partner(env: Env, name: str = "Partenaire Test") -> int:
    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, env.seeded.agency_a)
        manager = _user(s, env.seeded.manager_a)
        partner = partners.create_partner(s, manager, name, now=T0)
        partners.set_fee(s, partner, Decimal("300"), date(2026, 1, 1), manager, now=T0)
        s.commit()
        return partner.id


def add_sale(env: Env, pid: int, *tickets: TicketInput) -> int:
    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, env.seeded.agency_a)
        result = sales.create_sale(
            s, _user(s, env.seeded.manager_a), partner_sale(pid, *tickets), day=OCT_3, now=T0
        )
        s.commit()
        return result.sale.id


def add_line(
    env: Env, pid: int, kind: str, amount: str, reason: str, day: str, *, by: int | None = None,
    direction: str = "",
) -> int:  # fmt: skip
    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, env.seeded.agency_a)
        written = partner_book.add_line(
            s, _user(s, by or env.seeded.manager_a), pid, kind=kind, amount=amount,
            reason=reason, direction=direction, request_key=partner_book.new_request_key(),
            entry_date=day, now=T0,
        )  # fmt: skip
        s.commit()
        entry_id: int = written.row.id
        return entry_id


def add_payment(env: Env, pid: int, amount: str, day: str) -> None:
    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, env.seeded.agency_a)
        partner_book.record_payment(
            s, _user(s, env.seeded.manager_a), pid, amount=amount, method="cash",
            received_by_user_id=env.seeded.manager_a,
            request_key=partner_book.new_request_key(), paid_on=day, now=T0,
        )  # fmt: skip
        s.commit()


def close_at(env: Env, pid: int, start: date, now: datetime) -> None:
    set_days(env, "closed")
    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, env.seeded.agency_a)
        partner_book.close_period(
            s, _user(s, env.seeded.manager_a), pid, start,
            request_key=partner_book.new_request_key(), now=now,
        )  # fmt: skip
        s.commit()


def statement_rows(env: Env, pid: int, start: date, today: date | None = None) -> list[Any]:
    from app import partner_statement

    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, env.seeded.agency_a)
        statement = partner_statement.statement_for(
            s, _user(s, env.seeded.manager_a), pid, start, today=today or T0.date()
        )
    return rows(workbook(partner_statement.build_statement_xlsx(statement)))


def from_tickets(found: list[Any]) -> list[Any]:
    return found[found.index(TICKET_HEAD) :]


NOV_5 = datetime(2026, 11, 5, 9, 0, tzinfo=T0.tzinfo)


def test_several_tickets_an_unknown_airline_a_refund_and_a_rejected_line(env: Env) -> None:
    acc_id = add_user(
        env.db.owner_url, env.seeded.agency_a, ACCOUNTANT_PHONE, role=ROLE_ACCOUNTANT,
        display_name="Comptable A",
    )  # fmt: skip
    pid = new_partner(env)
    unknown = TicketInput(
        doc_number="3940000003", airline="999", total_doc="5000", pax_name="SYNTH02/PAX",
        pax_type="ADT", rloc="XYZ789", trnc="TKTT",
    )  # fmt: skip
    add_sale(env, pid, ticket("3940000001", "10000"), unknown)
    add_line(env, pid, "RFND", "700", "remboursement billet", "2026-10-10")
    rejected = add_line(
        env, pid, "AJUSTEMENT", "200", "frais oubliés test", "2026-10-11", by=acc_id,
        direction="plus",
    )  # fmt: skip
    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, env.seeded.agency_a)
        partner_book.decide_line(
            s, _user(s, env.seeded.manager_a), pid, rejected, decision="reject",
            note="pas justifié ici", now=T0,
        )  # fmt: skip
        s.commit()
    close_first_period(env, pid)
    assert from_tickets(statement_rows(env, pid, OCT_1)) == [
        TICKET_HEAD,
        (dt(3), "3940000001", "Tunisair", "SYNTH01/PAX", "ABC123", D("10000"), D("300"),
         D("10300")),
        (dt(3), "3940000003", "Cie 999", "SYNTH02/PAX", "XYZ789", D("5000"), D("300"),
         D("5300")),
        ("Total des billets", None, None, None, None, D("15000"), D("600"), D("15600")),
        (),
        ("Lignes manuelles",),
        LINE_HEAD,
        (dt(10), "Remboursement (RFND)", "remboursement billet", None, None, None, "comptée",
         D("-700")),
        (dt(11), "Ajustement", "frais oubliés test", None, None, None,
         "rejetée (non comptée)", D("200")),
        total("Total des lignes comptées", "-700"),
        (),
        ("Paiements",),
        PAYMENT_HEAD,
        ("Aucun paiement.",),
        total("Total des paiements", "0"),
        (),
        ("Récapitulatif",),
        total("Solde précédent", "0"),
        total("Billets (+)", "15600"),
        total("Lignes manuelles (+/−)", "-700"),
        total("Paiements (−)", "0"),
        total("Nouveau solde", "14900"),
        ("Sens du solde", "Le partenaire doit ce montant à l'agence."),
        SIGN,
    ]  # fmt: skip


def test_a_payments_only_period_opens_on_the_closed_balance_and_ends_negative(env: Env) -> None:
    pid = new_partner(env)
    add_sale(env, pid, ticket("3940000001", "1000"))
    close_first_period(env, pid)  # 1 to 15 October: 1 300 owed
    add_payment(env, pid, "2000", "2026-10-17")
    close_at(env, pid, OCT_16, NOV_5)
    found = statement_rows(env, pid, OCT_16, NOV_5.date())
    assert ("État", "Clôturée le 05/11/2026") in found
    assert ("Période", "du 16/10/2026 au 31/10/2026") in found
    assert from_tickets(found) == [
        TICKET_HEAD,
        ("Aucun billet dans cette période.",),
        ("Total des billets", None, None, None, None, D("0"), D("0"), D("0")),
        (),
        ("Lignes manuelles",),
        LINE_HEAD,
        ("Aucune ligne.",),
        total("Total des lignes comptées", "0"),
        (),
        ("Paiements",),
        PAYMENT_HEAD,
        (dt(17), "Espèces", None, None, None, None, None, D("2000")),
        total("Total des paiements", "2000"),
        (),
        ("Récapitulatif",),
        total("Solde précédent", "1300"),
        total("Billets (+)", "0"),
        total("Lignes manuelles (+/−)", "0"),
        total("Paiements (−)", "2000"),
        total("Nouveau solde", "-700"),
        ("Sens du solde", "L'agence doit ce montant au partenaire."),
        SIGN,
    ]


def test_a_zero_balance_says_the_account_is_settled(env: Env) -> None:
    pid = new_partner(env)
    add_sale(env, pid, ticket("3940000001", "1000"))
    add_payment(env, pid, "1300", "2026-10-12")
    found = statement_rows(env, pid, OCT_1)
    assert found[found.index(("Récapitulatif",)) :] == [
        ("Récapitulatif",),
        total("Solde précédent", "0"),
        total("Billets (+)", "1300"),
        total("Lignes manuelles (+/−)", "0"),
        total("Paiements (−)", "1300"),
        total("Nouveau solde", "0"),
        ("Sens du solde", "Compte soldé."),
        SIGN,
    ]


def test_an_empty_period_has_every_section_and_its_empty_sentence(env: Env) -> None:
    pid = new_partner(env)
    add_sale(env, pid, ticket("3940000001", "1000"))
    add_payment(env, pid, "1300", "2026-10-12")
    close_first_period(env, pid)
    found = statement_rows(env, pid, OCT_16)
    assert ("RELEVÉ DE COMPTE PARTENAIRE — PROVISOIRE",) in found
    assert from_tickets(found) == [
        TICKET_HEAD,
        ("Aucun billet dans cette période.",),
        ("Total des billets", None, None, None, None, D("0"), D("0"), D("0")),
        (),
        ("Lignes manuelles",),
        LINE_HEAD,
        ("Aucune ligne.",),
        total("Total des lignes comptées", "0"),
        (),
        ("Paiements",),
        PAYMENT_HEAD,
        ("Aucun paiement.",),
        total("Total des paiements", "0"),
        (),
        ("Récapitulatif",),
        total("Solde précédent", "0"),
        total("Billets (+)", "0"),
        total("Lignes manuelles (+/−)", "0"),
        total("Paiements (−)", "0"),
        total("Nouveau solde", "0"),
        ("Sens du solde", "Compte soldé."),
        SIGN,
    ]


def test_a_partner_name_starting_with_equals_is_an_exact_text_cell(env: Env) -> None:
    save_details(env)
    pid, _ = setup_account(env, name="=SOMME(1;2) Partenaire")
    close_first_period(env, pid)
    ws = workbook(env.manager_a().get(url(pid)).content)
    assert ws["A10"].value == "Partenaire"
    assert ws["B10"].value == "'=SOMME(1;2) Partenaire" and ws["B10"].data_type == "s"
    assert rows(ws) == closed_statement_rows("'=SOMME(1;2) Partenaire")


# --- invisible characters -------------------------------------------------------------------------


@pytest.mark.parametrize(
    "text",
    ["　=1+1", "﻿=1+1", "​=1+1", " @SUM(A1)", " +1", " -1",
     " 　﻿=1"],
)  # fmt: skip
def test_neutralise_sees_a_formula_behind_wide_or_invisible_blanks(text: str) -> None:
    from app.export_apres import neutralise

    assert neutralise(text) == "'" + text


def _statement(**overrides: Any) -> Any:
    from app import partner_statement as ps

    zero = Decimal("0.00")
    fields: dict[str, Any] = {
        "header": ("Agence Test",),
        "currency": "MRU",
        "partner_id": 1,
        "partner_name": "Partenaire Test",
        "start": OCT_1,
        "end": date(2026, 10, 15),
        "provisional": False,
        "closed_on": date(2026, 10, 18),
        "issued_on": date(2026, 10, 18),
        "period_id": 1,
        "body": ps.Body((), (), (), zero, zero, zero, zero, zero),
    }
    fields.update(overrides)
    return ps.Statement(**fields)


def test_bidi_and_zero_width_characters_are_removed_from_every_text_cell() -> None:
    from app import partner_statement as ps

    body = ps.Body(
        (ps.TicketRow(OCT_3, "394​0000001", "Tunisair", "SYN‮TH01/PAX", "AB﻿C",
                      D("1.00"), D("0.00"), D("1.00")),),
        (ps.EntryRow(OCT_3, "ADM (pénalité de la compagnie)", '　=HYPERLINK("x")',
                     ps.COUNTED, D("1.00")),),
        (ps.PaymentRow(OCT_3, "Espèces", "‮​+1+1", D("2.00")),),
        D("0.00"), D("1.00"), D("1.00"), D("2.00"), D("0.00"),
    )  # fmt: skip
    statement = _statement(
        partner_name="　​=Partenaire", header=("Agence‮ Test", "﻿@SUM(1)"),
        body=body,
    )  # fmt: skip
    ws = workbook(ps.build_statement_xlsx(statement))
    texts = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    for bad in ("‮", "​", "﻿", "　"):
        assert not any(bad in t for t in texts), repr(bad)
    for expected in (
        "'=Partenaire", "Agence Test", "'@SUM(1)", "3940000001", "SYNTH01/PAX", "ABC",
        "'=HYPERLINK(\"x\")", "'+1+1",
    ):  # fmt: skip
        assert expected in texts, expected
    assert [c.coordinate for row in ws.iter_rows() for c in row if c.data_type == "f"] == []


# --- a closed statement is never provisional ------------------------------------------------------


def test_a_closed_statement_without_its_close_date_is_refused() -> None:
    from app import partner_statement as ps

    with pytest.raises(ps.StatementError):
        ps.build_statement_xlsx(_statement(provisional=False, closed_on=None))
    row = PartnerPeriod(
        agency_id=1, partner_id=1, start_date=OCT_1, end_date=date(2026, 10, 15),
        status="closed", closed_at=None,
    )  # fmt: skip
    with pytest.raises(ps.StatementError):
        ps.closed_date(row)
    found = rows(workbook(ps.build_statement_xlsx(_statement())))
    assert ("État", "Clôturée le 18/10/2026") in found
    assert not any("PROVISOIRE" in str(v) for r in found for v in r)


# --- one snapshot read ----------------------------------------------------------------------------


def test_only_the_snapshot_of_the_wanted_period_is_read(env: Env) -> None:
    from sqlalchemy import event

    from app import partner_statement

    pid = new_partner(env)
    add_sale(env, pid, ticket("3940000001", "1000"))
    close_first_period(env, pid)
    add_payment(env, pid, "500", "2026-10-17")
    close_at(env, pid, OCT_16, NOV_5)
    seen: list[str] = []

    def listener(*args: Any) -> None:
        seen.append(args[2])

    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, env.seeded.agency_a)
        engine = s.get_bind()
        event.listen(engine, "before_cursor_execute", listener)
        try:
            statement = partner_statement.statement_for(
                s, _user(s, env.seeded.manager_a), pid, OCT_1, today=NOV_5.date()
            )
        finally:
            event.remove(engine, "before_cursor_execute", listener)
    assert statement.closing == Decimal("1300.00") and not statement.provisional
    reads = [q for q in seen if "partner_periods.snapshot" in q]
    assert len(reads) == 1, reads
    assert "partner_periods.id =" in reads[0]


# --- a corrupt snapshot is a 409, never a 500 -----------------------------------------------------


def _period_id(env: Env, pid: int) -> int:
    with owner_session(env.db.owner_url) as s:
        rebind_tenant(s, env.seeded.agency_a)
        found: int = s.scalars(
            select(PartnerPeriod.id).where(PartnerPeriod.partner_id == pid)
        ).one()
        return found


def test_a_corrupt_snapshot_is_a_409_page_and_one_log_line(
    env: Env, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app import partner_statement
    from app.web.routes import partner_statements as routes

    pid, _ = setup_account(env, name="Partenaire Secret")
    close_first_period(env, pid)

    def corrupt(*args: Any, **kwargs: Any) -> Any:
        raise partner_statement.StatementError(partner_statement.BAD_SNAPSHOT)

    logged: list[tuple[str, dict[str, Any]]] = []

    class FakeLog:
        def error(self, message: str, *, extra: dict[str, Any]) -> None:
            logged.append((message, extra))

    monkeypatch.setattr(partner_statement, "from_snapshot", corrupt)
    monkeypatch.setattr(routes, "log", FakeLog())
    response = env.manager_a().get(url(pid))
    assert response.status_code == 409
    assert response.headers["content-type"].startswith("text/html")
    assert response.headers["cache-control"] == "no-store"
    assert "Le relevé figé de cette période est illisible" in response.text
    [(message, extra)] = logged
    assert extra == {"partner_id": pid, "period_id": _period_id(env, pid)}
    assert "Secret" not in message
    assert audit_events(env.db.owner_url, env.seeded.agency_a, "partner_statement_downloaded") == []

"""Only the manager sees the profit (owner's decision, 28 Oct 2026, docs/design/02).

The profit (P.B minus the cost) and the TJQ APRES margin row are shown to the manager and to
nobody else: the accountant and the agents get no « Profit » word and no profit figure in any
page, fragment, chat answer or downloaded workbook. P.B, the cost and the partner fee stay where
they were. Synthetic data only.
"""

from __future__ import annotations

import io
import json
from collections.abc import Iterator
from datetime import UTC, datetime
from decimal import Decimal
from typing import Any

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import apres_compare, export_apres, permissions
from app.chat import tools
from app.chat.reply import Reply
from app.chat.tools import run_tool
from app.ingest import ingest_report
from app.models import ROLE_ACCOUNTANT, ROLE_MANAGER, ROLE_STAFF, DayExport, User
from app.parsing.amadeus import parse_apres
from tests.chat_support import ChatWorld
from tests.fixtures import synthetic_20oct as fx
from tests.page_support import Env, fr
from tests.test_sale_pages import (
    SIGN,
    TODAY,
    accountant,
    enter,
    form,
    make_partner,
    post,
    table,
    text_of,
    tk,
)
from tests.web_support import WebDb, owner_session

pytest_plugins = ["tests.chat_fixtures"]  # the fixtures ``w`` and ``chat_template``

NNBSP = " "  # the narrow no-break space the pages put between thousands
COST, PB = "21368.40", "24000"
PROFIT = f"2{NNBSP}631,60"  # 24000 - 21368.40, as the pages print it
ROLES = ("manager", "accountant", "staff")
PROFIT_WORDS = ("profit", "marge", "margin", "bénéfice", "benefice")


@pytest.fixture
def env(web_db: WebDb) -> Env:
    e = Env(web_db)
    e.link(e.seeded.staff_a, SIGN)
    return e


def clients(env: Env) -> dict[str, TestClient]:
    return {"manager": env.manager_a(), "accountant": accountant(env), "staff": env.staff_a()}


def sale_fields() -> dict[str, Any]:
    return {"staff_sign": SIGN, "day": TODAY, "pb": PB}


def assert_profit_shown(html: str) -> None:
    text = text_of(html)
    assert "Profit" in text and PROFIT in text


def assert_no_profit(html: str) -> None:
    text = text_of(html)
    assert "Profit" not in html and "profit" not in html.lower()
    assert PROFIT not in text and "2631" not in html and "2 631" not in text


# --- the permission -----------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("role", "expected"), [(ROLE_MANAGER, True), (ROLE_ACCOUNTANT, False), (ROLE_STAFF, False)]
)
def test_only_the_manager_may_see_the_profit(role: str, expected: bool) -> None:
    assert permissions.can_see_profit(User(agency_id=1, phone="+1", role=role)) is expected


# --- the pages ----------------------------------------------------------------------------------


@pytest.mark.parametrize("role", ROLES)
def test_the_read_back_shows_the_profit_to_the_manager_only(env: Env, role: str) -> None:
    c = clients(env)[role]
    r = post(c, "/sales/preview", form(tk(cost=COST), **sale_fields()))
    assert r.status_code == 200, r.text[:2000]
    assert "24" in text_of(r.text)  # P.B and cost are still there for everybody
    assert f"21{NNBSP}368,40" in text_of(r.text)
    if role == "manager":
        assert_profit_shown(r.text)
    else:
        assert_no_profit(r.text)


@pytest.mark.parametrize("role", ROLES)
def test_the_sale_page_shows_the_profit_to_the_manager_only(env: Env, role: str) -> None:
    c = clients(env)
    sale_id = enter(c["staff"], tk(cost=COST), **sale_fields())
    r = c[role].get(f"/sales/{sale_id}")
    assert r.status_code == 200
    assert f"21{NNBSP}368,40" in text_of(r.text) and fr(24000) in text_of(r.text)
    if role == "manager":
        assert_profit_shown(r.text)
    else:
        assert_no_profit(r.text)


@pytest.mark.parametrize("role", ROLES)
def test_the_sales_list_shows_the_profit_to_the_manager_only(env: Env, role: str) -> None:
    c = clients(env)
    enter(c["staff"], tk(cost=COST), **sale_fields())
    r = c[role].get(f"/sales?date={TODAY}")
    assert r.status_code == 200
    assert f"21{NNBSP}368,40" in text_of(r.text) and fr(24000) in text_of(r.text)
    if role == "manager":
        assert text_of(r.text).count(PROFIT) == 2  # the row and the day's total
        assert "Profit" in r.text
    else:
        assert_no_profit(r.text)


@pytest.mark.parametrize("role", ROLES)
def test_the_day_list_shows_the_profit_to_the_manager_only(env: Env, role: str) -> None:
    c = clients(env)
    enter(c["staff"], tk(cost=COST), **sale_fields())
    r = c[role].get("/days")
    assert r.status_code == 200
    # the day's cost and P.B: for manager and accountant only (decision 26)
    shown = f"21{NNBSP}368,40" in text_of(r.text) and fr(24000) in text_of(r.text)
    assert shown is (role != "staff")
    if role == "manager":
        assert_profit_shown(r.text)
    else:
        assert_no_profit(r.text)


@pytest.mark.parametrize("role", ROLES)
def test_the_day_page_shows_the_profit_to_the_manager_only(env: Env, role: str) -> None:
    c = clients(env)
    enter(c["staff"], tk(cost=COST), **sale_fields())
    r = c[role].get(f"/days/{TODAY}")
    assert r.status_code == 200
    assert f"21{NNBSP}368,40" in text_of(r.text) and fr(24000) in text_of(r.text)
    if role == "manager":
        assert text_of(r.text).count(PROFIT) == 2  # the day's total and the sign's row
    else:
        assert_no_profit(r.text)


@pytest.mark.parametrize("role", ["manager", "accountant"])
def test_the_partner_fee_stays_visible_to_manager_and_accountant(env: Env, role: str) -> None:
    pid = make_partner(env, "BADI", "500")
    c = clients(env)
    sale_id = enter(c["manager"], tk(cost="15000"), staff_sign=SIGN, day=TODAY, ag=str(pid))
    text = text_of(c[role].get(f"/sales/{sale_id}").text)
    assert "Frais du partenaire" in text and fr(500) in text
    assert ("Profit" in text) is (role == "manager")


@pytest.mark.parametrize("role", ["accountant", "staff"])
def test_the_home_page_has_no_profit_either(env: Env, role: str) -> None:
    c = clients(env)
    enter(c["staff"], tk(cost=COST), **sale_fields())
    r = c[role].get("/")
    assert r.status_code == 200
    assert_no_profit(r.text)


# --- the chat -----------------------------------------------------------------------------------


def _blob(answer: Reply) -> str:
    """Everything a chat answer carries: its lines (shown, and sent to the model) and its data."""
    lines = " ".join(f"{line.text} {line.href or ''}" for line in answer.lines)
    return f"{lines} {json.dumps(answer.data, default=str, ensure_ascii=False)}".lower()


@pytest.mark.parametrize("role", ROLES)
def test_no_chat_answer_carries_a_profit(w: ChatWorld, role: str) -> None:
    """No lookup returns a profit or a margin to anybody (the chat never had one); for the roles
    that may not see it, it stays that way in what is shown and in what reaches the model."""
    user = {"manager": w.manager, "accountant": w.accountant, "staff": w.staff}[role]
    s = w.scenario
    calls: list[tuple[str, dict[str, Any]]] = [
        ("list_sales", {"day": "2026-10-20"}),
        ("list_sales", {}),
        ("ticket_status", {"doc": s.complete}),
        ("ticket_status", {"doc": s.different}),
        ("open_flags", {}),
        ("not_sent_tickets", {}),
        ("day_sheet", {"day": "2026-10-20"}),
        ("void_link", {"doc": s.plain}),
        ("help", {}),
    ]
    for name, args in calls:
        answer = run_tool(w.ctx(user), name, args)
        blob = _blob(answer)
        for word in PROFIT_WORDS:
            assert word not in blob, (name, word)
    for t in tools.tools_for(user):
        assert not any(word in t.description.lower() for word in PROFIT_WORDS)


# --- the TJQ APRES ------------------------------------------------------------------------------


D20 = "2026-10-20"
FIRST_ROW = 6
MARGIN_LABEL = "P.B − TOTAL DOC"


def _upload_20oct(env: Env) -> None:
    env.clock.now = datetime(2026, 10, 21, 12, 0, tzinfo=UTC)
    with owner_session(env.db.owner_url) as s:
        ingest_report(
            s, env.seeded.agency_a, fx.build_synthetic_20oct(), "20OCT.xlsx",
            uploaded_by=env.seeded.manager_a,
        )  # fmt: skip


def _sell_line(c: TestClient, index: int) -> None:
    line = [x for x in fx.LINES if x["TRNC"] == "TKTT"][index]
    ticket = tk(
        line["DOC NUMBER"], str(line["TOTAL DOC"]), airline=line["A/L"], pax_name=line["PAX NAME"],
        rloc=line["RLOC"],
    )  # fmt: skip
    enter(c, ticket, staff_sign=line["AS"], day=D20, pb=str(line["TOTAL DOC"] + 1000))


def _sheet(data: bytes) -> Any:
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert ws is not None
    return ws


def _totals_row(ws: Any) -> int:
    rows = [r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "TOTAUX"]
    assert len(rows) == 1
    return rows[0]


def _cells(ws: Any) -> list[Any]:
    return [c.value for row in ws.iter_rows() for c in row]


def _has_margin(data: bytes) -> bool:
    ws = _sheet(data)
    row = _totals_row(ws)
    label, value = ws.cell(row=row, column=10).value, ws.cell(row=row, column=11).value
    if label is None and value is None:
        assert MARGIN_LABEL not in _cells(ws)
        return False
    assert label == MARGIN_LABEL and value is not None
    return True


@pytest.fixture
def closed(env: Env, gate_open: None) -> Iterator[Env]:
    _upload_20oct(env)
    m = env.manager_a()
    _sell_line(m, 0)
    _sell_line(m, 1)
    assert post(m, f"/days/{D20}/close", {}, via=f"/days/{D20}/close").status_code == 303
    yield env


def _stored(env: Env) -> list[DayExport]:
    return sorted(table(env, DayExport), key=lambda e: e.version)


def test_the_stored_export_keeps_its_margin_row_unchanged(closed: Env) -> None:
    (stored,) = _stored(closed)
    assert _has_margin(stored.content)


def test_the_manager_downloads_the_stored_file_as_it_is(closed: Env) -> None:
    (stored,) = _stored(closed)
    got = closed.manager_a().get(f"/days/{D20}/apres.xlsx")
    assert got.status_code == 200 and got.content == stored.content
    assert _has_margin(got.content)


def test_the_accountants_download_has_no_margin_row(closed: Env) -> None:
    (stored,) = _stored(closed)
    acc = accountant(closed)
    got = acc.get(f"/days/{D20}/apres.xlsx")
    assert got.status_code == 200 and got.content[:2] == b"PK"
    assert not _has_margin(got.content)
    # everything else is the stored file: headers, lines, the other totals, the notes
    mine, ref = _sheet(got.content), _sheet(stored.content)
    total = _totals_row(ref)
    assert _totals_row(mine) == total
    for row in range(1, total + 1):
        for col in range(1, 14):
            if row == total and col in (10, 11):
                continue
            a, b = mine.cell(row=row, column=col), ref.cell(row=row, column=col)
            assert a.value == b.value, (row, col)
            assert (a.comment.text if a.comment else None) == (
                b.comment.text if b.comment else None
            )
    # the same version gives the same bytes every time, and ?v=1 too
    assert acc.get(f"/days/{D20}/apres.xlsx").content == got.content
    assert acc.get(f"/days/{D20}/apres.xlsx?v=1").content == got.content
    # the stored row is untouched
    (again,) = _stored(closed)
    assert again.content == stored.content and _has_margin(again.content)


def test_the_accountants_file_compares_exactly_like_the_stored_one(closed: Env) -> None:
    (stored,) = _stored(closed)
    got = accountant(closed).get(f"/days/{D20}/apres.xlsx").content
    ours, hers = parse_apres(stored.content), parse_apres(got)
    assert ours.tickets == hers.tickets
    result = apres_compare.compare(ours, hers)
    assert result.same and result.identical == result.matched == len(ours.tickets)
    assert export_apres.HEADERS == (
        "SEQ NO", "A/L", "DOC NUMBER", "TOTAL DOC", "PAX NAME", "AS", "TRNC", "RLOC",
        "P.B", "AG", "CLT", "AP",
    )  # fmt: skip


@pytest.mark.parametrize("role", ["manager", "accountant"])
def test_the_draft_has_the_margin_row_for_the_manager_only(
    env: Env, gate_open: None, role: str
) -> None:
    _upload_20oct(env)
    m = env.manager_a()
    _sell_line(m, 0)
    c = m if role == "manager" else accountant(env)
    got = c.get(f"/days/{D20}/apres-preview.xlsx")
    assert got.status_code == 200
    assert _has_margin(got.content) is (role == "manager")


def test_the_workbook_builder_can_leave_the_margin_out() -> None:
    from tests.test_export_apres import hostile_report

    report = hostile_report("x")
    assert _has_margin(export_apres.build_apres_xlsx(report))
    plain = export_apres.build_apres_xlsx(report, with_margin=False)
    assert not _has_margin(plain)


def test_removing_the_margin_from_a_stored_file_is_stable() -> None:
    from tests.test_export_apres import hostile_report

    full = export_apres.build_apres_xlsx(hostile_report("x"))
    once = export_apres.without_margin(full)
    assert not _has_margin(once)
    assert export_apres.without_margin(full) == once
    assert export_apres.without_margin(once) == export_apres.without_margin(once)
    assert Decimal(str(_sheet(once).cell(row=_totals_row(_sheet(once)), column=9).value)) == (
        Decimal(str(_sheet(full).cell(row=_totals_row(_sheet(full)), column=9).value))
    )


def test_stripping_the_margin_fails_closed_without_a_totals_row() -> None:
    # A workbook with no « TOTAUX » row: refuse rather than send a margin left in place.
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.cell(row=1, column=1, value="TJQ APRES")
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="totals row"):
        export_apres.without_margin(buf.getvalue())
